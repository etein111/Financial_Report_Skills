"""content_inserter.py - 内容插入模块

提供 abstract、english_summary 等内容插入到文档的功能。
"""

from typing import Dict, Any
from lxml import etree

try:
    from docx import Document
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False
    Document = None

import common
W = common.W_NS


def _find_marker_in_cells(table, marker: str):
    """在表格单元格中查找 marker。

    Args:
        table: 表格对象
        marker: 要查找的标记文本

    Returns:
        找到的单元格，或 None
    """
    for row in table.rows:
        for cell in row.cells:
            if marker in cell.text:
                return cell
    return None


def insert_abstract_to_cell(doc: Document, abstract_doc: Document, marker: str) -> bool:
    """在文档的表格单元格中插入 abstract 内容。

    查找包含 marker 的表格单元格，将 abstract_doc 的内容插入其中。
    使用 deepcopy 保留原始段落的完整格式（加粗、颜色等）。

    Args:
        doc: 目标 Document 对象
        abstract_doc: 包含 abstract 内容的 Document 对象
        marker: 单元格中的标记文本

    Returns:
        是否成功插入
    """
    if not DOCX_AVAILABLE:
        return False

    try:
        from copy import deepcopy

        # 在表格中查找 marker
        for table in doc.tables:
            cell = _find_marker_in_cells(table, marker)
            if cell:
                # 获取 abstract 段落 XML 元素（保留完整 XML）
                abstract_para_elems = [p._element for p in abstract_doc.paragraphs if p.text.strip()]

                if abstract_para_elems:
                    # 获取单元格内所有段落元素
                    cell_para_elems = [p._element for p in cell.paragraphs]

                    if cell_para_elems:
                        # 在第一个单元格段落之前插入 abstract 段落
                        first_cell_para = cell_para_elems[0]
                        parent = first_cell_para.getparent()

                        if parent is not None:
                            insert_idx = list(parent).index(first_cell_para)
                            for i, abstract_para_elem in enumerate(abstract_para_elems):
                                parent.insert(insert_idx + i, deepcopy(abstract_para_elem))

                            # 删除原始段落
                            for para_elem in cell_para_elems:
                                try:
                                    parent.remove(para_elem)
                                except ValueError:
                                    pass

                return True

        return False

    except Exception as e:
        print(f"  插入 abstract 到表格单元格失败: {e}")
        import traceback
        traceback.print_exc()
        return False


def _insert_abstract_at_paragraph(doc: Document, abstract_doc: Document, marker: str) -> bool:
    """在文档的段落中插入 abstract 内容。

    查找包含 marker 的段落，用 abstract_doc 的内容替换该段落。

    Args:
        doc: Document 对象
        abstract_doc: 包含 abstract 内容的 Document 对象
        marker: 段落中的标记文本

    Returns:
        是否成功插入
    """
    if not DOCX_AVAILABLE:
        return False

    try:
        target_para = None
        target_idx = None

        all_paras = list(doc.element.body.iter(f'{W}p'))
        for idx, para in enumerate(all_paras):
            text = common.get_element_text(para)
            if marker in text:
                target_para = para
                target_idx = idx
                break

        if target_para is None:
            return False

        # 获取 abstract 内容
        abstract_paras = [p for p in abstract_doc.paragraphs if p.text.strip()]

        # 获取 target_para 的父元素
        body = doc.element.body
        parent = target_para.getparent()

        if parent is None:
            return False

        # 在 target_para 之前插入新的段落
        insert_idx = list(parent).index(target_para)

        for i, abstract_para in enumerate(abstract_paras):
            new_para = abstract_para._element
            parent.insert(insert_idx + i, new_para)

        # 删除原始的 marker 段落
        parent.remove(target_para)

        return True

    except Exception as e:
        print(f"  在段落中插入 abstract 失败: {e}")
        return False


def _apply_english_summary_margins_doc(doc: Document, margins: Dict[str, str]) -> bool:
    """应用英文摘要部分的页边距（直接操作 Document 对象）。

    找到英文摘要内容之前最近的分节符，应用指定的页边距。

    Args:
        doc: Document 对象
        margins: 页边距字典

    Returns:
        是否成功
    """
    if not margins:
        return False

    body = doc.element.body
    all_elements = list(body)

    es_start_idx = None
    for idx, elem in enumerate(all_elements):
        if elem.tag == f'{W}p':
            text = common.get_element_text(elem).strip()
            if text == common.APPENDIX_1_MARKER:
                es_start_idx = idx
                break

    if es_start_idx is None or es_start_idx == 0:
        return False

    for i in range(es_start_idx - 1, -1, -1):
        elem = all_elements[i]
        if elem.tag == f'{W}p':
            pPr = elem.find(f'{W}pPr')
            if pPr is not None:
                sectPr = pPr.find(f'{W}sectPr')
                if sectPr is not None:
                    pgMar = sectPr.find(f'{W}pgMar')
                    if pgMar is None:
                        pgMar = etree.SubElement(sectPr, f'{W}pgMar')
                    for attr, val in margins.items():
                        pgMar.set(f'{W}{attr}', val)
                    print(f"  已应用 english_summary 页边距")
                    return True
    return False


