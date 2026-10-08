"""test_placeholder_fixer.py - placeholder_fixer.py 模块测试。

测试占位符修复功能。
"""

import pytest
from pathlib import Path
import zipfile

class TestFixRunsInParent:
    """测试 fix_runs_in_parent 函数。"""

    def test_handles_empty_collected_runs(self):
        """测试：处理空的 collected_runs。"""
        from placeholder_fixer import fix_runs_in_parent

        # 不应该抛出异常
        fix_runs_in_parent(None, [])


class TestPlaceholderConstants:
    """测试占位符常量。"""

    def test_placeholder_markers_defined(self):
        """测试：占位符标记已定义。"""
        from placeholder_fixer import PLACEHOLDER_START, PLACEHOLDER_END

        assert PLACEHOLDER_START == '{{'
        assert PLACEHOLDER_END == '}}'


class TestFixSplitPlaceholdersInDocx:
    """测试 _fix_split_placeholders_in_docx 函数。"""

    def test_handles_nonexistent_file(self):
        """测试：处理不存在的文件。"""
        from placeholder_fixer import _fix_split_placeholders_in_docx

        # 不存在文件时应该返回原路径而不是崩溃
        result = _fix_split_placeholders_in_docx("nonexistent_file.docx")

        assert result == "nonexistent_file.docx"

    def test_handles_file_without_document_xml(self, tmp_path):
        """测试：处理没有 document.xml 的文件。"""
        from placeholder_fixer import _fix_split_placeholders_in_docx

        # 创建一个只有部分结构的 zip
        doc_path = tmp_path / "incomplete.docx"
        with zipfile.ZipFile(str(doc_path), 'w') as z:
            z.writestr("[Content_Types].xml", b"<Types/>")

        result = _fix_split_placeholders_in_docx(str(doc_path))

        # 应该返回原路径而不是崩溃
        assert result == str(doc_path)

    def test_handles_docx_without_body(self, tmp_path):
        """测试：处理没有 body 的 document.xml。"""
        from placeholder_fixer import _fix_split_placeholders_in_docx

        doc_path = tmp_path / "no_body.docx"
        with zipfile.ZipFile(str(doc_path), 'w') as z:
            z.writestr("[Content_Types].xml", b"<Types/>")
            z.writestr("word/document.xml", b"<document/>")

        result = _fix_split_placeholders_in_docx(str(doc_path))

        assert result == str(doc_path)

    def test_overwrites_file_when_output_is_none(self, tmp_path):
        """测试：output_path 为 None 时覆盖原文件。"""
        from placeholder_fixer import _fix_split_placeholders_in_docx

        try:
            from docx import Document
        except ImportError:
            pytest.skip("python-docx not installed")

        doc = Document()
        doc.add_paragraph("Simple document")

        doc_path = tmp_path / "to_overwrite.docx"
        doc.save(str(doc_path))

        result = _fix_split_placeholders_in_docx(str(doc_path), output_path=None)

        assert result == str(doc_path)
