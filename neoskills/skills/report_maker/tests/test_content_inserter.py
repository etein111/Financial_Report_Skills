"""test_content_inserter.py - content_inserter.py 模块测试。

测试内容插入到文档的功能。
"""

import pytest
from pathlib import Path
import shutil
from unittest.mock import MagicMock, patch

class TestFindMarkerInCells:
    """测试 _find_marker_in_cells 函数。"""

    def test_raises_error_when_table_is_none(self):
        """测试：table 为 None 时抛出异常。"""
        from content_inserter import _find_marker_in_cells

        with pytest.raises((AttributeError, TypeError)):
            _find_marker_in_cells(None, "test_marker")

    def test_finds_marker_in_cell(self):
        """测试：在单元格中找到 marker。"""
        from content_inserter import _find_marker_in_cells

        try:
            from docx import Document
        except ImportError:
            pytest.skip("python-docx not installed")

        doc = Document()
        table = doc.add_table(rows=1, cols=1)
        table.rows[0].cells[0].text = "Test {{MARKER}} content"

        result = _find_marker_in_cells(table, "{{MARKER}}")

        assert result is not None


class TestInsertAbstractToCell:
    """测试 insert_abstract_to_cell 函数。"""

    def test_returns_false_when_docx_unavailable(self):
        """测试：docx 不可用时返回 False。"""
        from content_inserter import insert_abstract_to_cell

        with patch('content_inserter.DOCX_AVAILABLE', False):
            result = insert_abstract_to_cell(None, None, "marker")

            assert result is False

    def test_returns_false_when_no_table_has_marker(self):
        """测试：表格中没有 marker 时返回 False。"""
        from content_inserter import insert_abstract_to_cell

        try:
            from docx import Document
        except ImportError:
            pytest.skip("python-docx not installed")

        doc = Document()
        doc.add_table(rows=2, cols=2)
        doc.add_paragraph("No marker here")

        abstract_doc = Document()
        abstract_doc.add_paragraph("Abstract content")

        result = insert_abstract_to_cell(doc, abstract_doc, "{{ABSTRACT}}")

        assert result is False

    def test_inserts_abstract_to_cell_with_marker(self, tmp_path, fixtures_dir):
        """测试：在带 marker 的单元格中插入 abstract。"""
        from content_inserter import insert_abstract_to_cell

        try:
            from docx import Document
        except ImportError:
            pytest.skip("python-docx not installed")

        doc_path = fixtures_dir / "doc_with_cell_marker.docx"
        abstract_path = fixtures_dir / "abstract_content.docx"

        # 复制到临时文件
        temp_doc = tmp_path / "test_cell.docx"
        temp_abstract = tmp_path / "abstract.docx"
        shutil.copy(doc_path, temp_doc)
        shutil.copy(abstract_path, temp_abstract)

        doc = Document(str(temp_doc))
        abstract_doc = Document(str(temp_abstract))

        result = insert_abstract_to_cell(doc, abstract_doc, "{{ABSTRACT}}")

        assert result is True


class TestInsertAbstractAtParagraph:
    """测试 _insert_abstract_at_paragraph 函数。"""

    def test_returns_false_when_docx_unavailable(self):
        """测试：docx 不可用时返回 False。"""
        from content_inserter import _insert_abstract_at_paragraph

        with patch('content_inserter.DOCX_AVAILABLE', False):
            result = _insert_abstract_at_paragraph(None, None, "marker")

            assert result is False

    def test_returns_false_when_marker_not_found(self):
        """测试：未找到 marker 时返回 False。"""
        from content_inserter import _insert_abstract_at_paragraph

        try:
            from docx import Document
        except ImportError:
            pytest.skip("python-docx not installed")

        doc = Document()
        doc.add_paragraph("No marker here")

        abstract_doc = Document()
        abstract_doc.add_paragraph("Abstract content")

        result = _insert_abstract_at_paragraph(doc, abstract_doc, "{{ABSTRACT}}")

        assert result is False

    def test_inserts_abstract_at_paragraph_with_marker(self, tmp_path, fixtures_dir):
        """测试：在带 marker 的段落中插入 abstract。"""
        from content_inserter import _insert_abstract_at_paragraph

        try:
            from docx import Document
        except ImportError:
            pytest.skip("python-docx not installed")

        doc_path = fixtures_dir / "doc_with_paragraph_marker.docx"
        abstract_path = fixtures_dir / "abstract_content.docx"

        # 复制到临时文件
        temp_doc = tmp_path / "test_para.docx"
        temp_abstract = tmp_path / "abstract.docx"
        shutil.copy(doc_path, temp_doc)
        shutil.copy(abstract_path, temp_abstract)

        doc = Document(str(temp_doc))
        abstract_doc = Document(str(temp_abstract))

        result = _insert_abstract_at_paragraph(doc, abstract_doc, "{{ABSTRACT}}")

        assert result is True
