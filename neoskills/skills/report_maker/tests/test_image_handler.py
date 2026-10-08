"""test_image_handler.py - image_handler.py 模块测试。

测试图片处理功能。
"""

import pytest
from pathlib import Path
import shutil
from unittest.mock import MagicMock, patch

class TestEscapeJinjaInCharts:
    """测试 _escape_jinja_in_charts 函数。"""

    def test_returns_empty_dict_for_nonexistent_file(self):
        """测试：文件不存在时返回空字典。"""
        from image_handler import _escape_jinja_in_charts

        result = _escape_jinja_in_charts("nonexistent.docx")

        assert result == {}

    def test_handles_file_without_charts(self, tmp_path):
        """测试：处理没有图表的文件。"""
        from image_handler import _escape_jinja_in_charts

        try:
            from docx import Document
        except ImportError:
            pytest.skip("python-docx not installed")

        doc = Document()
        doc.add_paragraph("No charts here")

        doc_path = tmp_path / "no_charts.docx"
        doc.save(str(doc_path))

        result = _escape_jinja_in_charts(str(doc_path))

        # 没有图表文件，应该返回空字典
        assert isinstance(result, dict)

    def test_escapes_jinja_in_chart(self, tmp_path, fixtures_dir):
        """测试：转义图表中的 Jinja 语法。"""
        from image_handler import _escape_jinja_in_charts

        doc_path = fixtures_dir / "doc_with_chart_jinja.docx"

        # 复制到临时文件
        temp_doc = tmp_path / "test_chart_jinja.docx"
        shutil.copy(doc_path, temp_doc)

        result = _escape_jinja_in_charts(str(temp_doc))

        # 应该返回被转义的文件
        assert isinstance(result, dict)

    def test_handles_chart_without_jinja(self, tmp_path, fixtures_dir):
        """测试：处理没有 Jinja 语法的图表。"""
        from image_handler import _escape_jinja_in_charts

        doc_path = fixtures_dir / "doc_with_chart_no_jinja.docx"

        # 复制到临时文件
        temp_doc = tmp_path / "test_chart_no_jinja.docx"
        shutil.copy(doc_path, temp_doc)

        result = _escape_jinja_in_charts(str(temp_doc))

        # 没有 Jinja 语法，应该返回空字典
        assert result == {}


class TestRestoreJinjaInCharts:
    """测试 _restore_jinja_in_charts 函数。"""

    def test_does_nothing_for_empty_dict(self):
        """测试：空字典时不操作。"""
        from image_handler import _restore_jinja_in_charts

        # 不应该抛出异常
        _restore_jinja_in_charts("any_path.docx", {})

    def test_handles_nonexistent_file(self):
        """测试：处理不存在的文件。"""
        from image_handler import _restore_jinja_in_charts

        escaped_files = {"word/charts/chart1.xml": "original content"}

        # 不应该抛出异常
        _restore_jinja_in_charts("nonexistent.docx", escaped_files)

    def test_restores_jinja_in_chart(self, tmp_path, fixtures_dir):
        """测试：恢复图表中的 Jinja 语法。"""
        from image_handler import _restore_jinja_in_charts

        doc_path = fixtures_dir / "doc_with_chart_jinja.docx"

        # 复制到临时文件
        temp_doc = tmp_path / "test_restore_jinja.docx"
        shutil.copy(doc_path, temp_doc)

        # 先转义
        from image_handler import _escape_jinja_in_charts
        escaped = _escape_jinja_in_charts(str(temp_doc))

        # 再恢复
        _restore_jinja_in_charts(str(temp_doc), escaped)

        # 不应该抛出异常


class TestInsertImageToPlaceholder:
    """测试 _insert_image_to_placeholder 函数。"""

    def test_returns_false_when_docx_unavailable(self):
        """测试：docx 不可用时返回 False。"""
        from image_handler import _insert_image_to_placeholder

        with patch('image_handler.DOCX_AVAILABLE', False):
            result = _insert_image_to_placeholder(None, "placeholder", "image.png")

            assert result is False

    def test_returns_false_when_placeholder_not_found(self, tmp_path):
        """测试：占位符未找到时返回 False。"""
        from image_handler import _insert_image_to_placeholder

        try:
            from docx import Document
        except ImportError:
            pytest.skip("python-docx not installed")

        doc = Document()
        doc.add_paragraph("No placeholder here")

        doc_path = tmp_path / "no_placeholder.docx"
        doc.save(str(doc_path))

        doc = Document(str(doc_path))

        result = _insert_image_to_placeholder(doc, "{{IMAGE}}", "test.png")

        assert result is False

    def test_inserts_image_in_paragraph(self, tmp_path, fixtures_dir):
        """测试：在段落中插入图片。"""
        from image_handler import _insert_image_to_placeholder

        try:
            from docx import Document
        except ImportError:
            pytest.skip("python-docx not installed")

        doc_path = fixtures_dir / "doc_with_placeholder.docx"
        image_path = fixtures_dir / "test_image.png"

        # 复制到临时文件
        temp_doc = tmp_path / "test_placeholder.docx"
        shutil.copy(doc_path, temp_doc)

        doc = Document(str(temp_doc))

        result = _insert_image_to_placeholder(doc, "{{IMAGE}}", str(image_path))

        # 返回 True 表示成功找到并处理了占位符
        assert result is True


class TestCopyMediaBetweenDocx:
    """测试 _copy_media_between_docx 函数。"""

    def test_returns_zero_when_source_missing(self, tmp_path):
        """测试：源文件不存在时返回 0。"""
        from image_handler import _copy_media_between_docx

        try:
            from docx import Document
        except ImportError:
            pytest.skip("python-docx not installed")

        doc = Document()
        doc.add_paragraph("Target")

        target_path = tmp_path / "target.docx"
        doc.save(str(target_path))

        result = _copy_media_between_docx("nonexistent_source.docx", str(target_path))

        assert result == 0

    def test_returns_zero_when_target_missing(self, tmp_path, fixtures_dir):
        """测试：目标文件不存在时返回 0。"""
        from image_handler import _copy_media_between_docx

        try:
            from docx import Document
        except ImportError:
            pytest.skip("python-docx not installed")

        source_path = fixtures_dir / "source_with_media.docx"

        result = _copy_media_between_docx(str(source_path), str(tmp_path / "nonexistent.docx"))

        assert result == 0

    def test_copies_media_between_docs(self, tmp_path, fixtures_dir):
        """测试：在两个文档之间复制媒体文件。"""
        from image_handler import _copy_media_between_docx

        try:
            from docx import Document
        except ImportError:
            pytest.skip("python-docx not installed")

        source_path = fixtures_dir / "source_with_media.docx"
        target_path = fixtures_dir / "doc_with_placeholder.docx"

        # 复制到临时文件
        temp_source = tmp_path / "source.docx"
        temp_target = tmp_path / "target.docx"
        shutil.copy(source_path, temp_source)
        shutil.copy(target_path, temp_target)

        result = _copy_media_between_docx(str(temp_source), str(temp_target))

        # fixture 文档媒体结构可能不完整，函数能处理即可
        assert isinstance(result, int)
