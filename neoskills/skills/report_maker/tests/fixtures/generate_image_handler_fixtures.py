#!/usr/bin/env python3
"""generate_image_handler_fixtures.py - 生成 image_handler 测试用的 fixture。

用于测试 image_handler.py 中的：
- _escape_jinja_in_charts: 转义图表中的 Jinja2 语法
- _restore_jinja_in_charts: 恢复图表中的 Jinja2 语法
- _insert_image_to_placeholder: 在占位符位置插入图片
- _copy_media_between_docx: 复制媒体文件

包含 Jinja 语法的图表 XML 示例。
"""

import sys
from pathlib import Path
import zipfile

try:
    from docx import Document
except ImportError:
    print("python-docx not installed. Run: uv pip install python-docx")
    sys.exit(1)


# 包含 Jinja 语法的图表 XML
CHART_XML_WITH_JINJA = b'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<c:chartSpace xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart"
              xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
              xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
    <c:date1904 val="0"/>
    <c:lang val="zh-CN"/>
    <c:roundedCorners val="0"/>
    <c:chart>
        <c:title>
            <c:tx>
                <c:rich>
                    <a:bodyPr/>
                    <a:lstStyle/>
                    <a:p>
                        <a:pPr/>
                        <a:r>
                            <a:rPr lang="zh-CN" b="1"/>
                            <a:t>{{chart_title}}</a:t>
                        </a:r>
                    </a:p>
                </c:rich>
            </c:tx>
            <c:overlay val="0"/>
        </c:title>
        <c:autoTitleDeleted val="0"/>
        <c:plotArea>
            <c:layout/>
            <c:barChart>
                <c:barDir val="col"/>
                <c:grouping val="clustered"/>
                <c:varyColors val="0"/>
                <c:ser>
                    <c:idx val="0"/>
                    <c:order val="0"/>
                    <c:tx>
                        <c:v>{{series_name}}</c:v>
                    </c:tx>
                    <c:spPr>
                        <a:solidFill>
                            <a:srgbClr val="4472C4"/>
                        </a:solidFill>
                    </c:spPr>
                </c:ser>
                <c:dLbls>
                    <c:showLegendKey val="0"/>
                    <c:showVal val="0"/>
                    <c:showCatName val="0"/>
                    <c:showSerName val="0"/>
                    <c:showPercent val="0"/>
                    <c:showBubbleSize val="0"/>
                </c:dLbls>
                <c:gapWidth val="150"/>
                <c:axId val="1"/>
                <c:axId val="2"/>
            </c:barChart>
            <c:catAx>
                <c:axId val="1"/>
                <c:scaling>
                    <c:orientation val="minMax"/>
                </c:scaling>
                <c:delete val="0"/>
                <c:axPos val="b"/>
                <c:numFmt formatCode="General" sourceLinked="1"/>
                <c:majorTickMark val="out"/>
                <c:minorTickMark val="none"/>
                <c:tickLblPos val="nextTo"/>
            </c:catAx>
            <c:valAx>
                <c:axId val="2"/>
                <c:scaling>
                    <c:orientation val="minMax"/>
                </c:scaling>
                <c:delete val="0"/>
                <c:axPos val="l"/>
                <c:majorGridlines/>
            </c:valAx>
        </c:plotArea>
        <c:legend>
            <c:legendPos val="r"/>
            <c:overlay val="0"/>
        </c:legend>
        <c:plotVisOnly val="1"/>
        <c:dispBlanksAs val="gap"/>
    </c:chart>
    <c:printSettings>
        <c:headerFooter/>
        <c:pageMargins b="0.75" l="0.7" r="0.7" t="0.75" header="0.3" footer="0.3"/>
        <c:pageSetup/>
    </c:printSettings>
</c:chartSpace>
'''

# 不包含 Jinja 语法的图表 XML
CHART_XML_WITHOUT_JINJA = b'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<c:chartSpace xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart">
    <c:chart>
        <c:title>
            <c:tx>
                <c:v>Stock Rating Distribution</c:v>
            </c:tx>
        </c:title>
    </c:chart>
</c:chartSpace>
'''


def create_doc_with_chart_and_jinja(output_path: str):
    """创建带图表（包含 Jinja 语法）的 docx。"""
    # 先创建一个简单的 docx
    doc = Document()
    doc.add_heading("研究报告", 0)
    doc.add_paragraph("以下是评级图表：")

    # 保存 docx
    doc.save(output_path)

    # 向 docx 中添加图表文件
    with zipfile.ZipFile(output_path, 'a') as z:
        # 创建图表目录
        chart_xml = CHART_XML_WITH_JINJA
        z.writestr('word/charts/chart1.xml', chart_xml)

    print(f"已生成: {output_path}")


def create_doc_with_chart_without_jinja(output_path: str):
    """创建带图表（不包含 Jinja 语法）的 docx。"""
    doc = Document()
    doc.add_heading("研究报告", 0)
    doc.add_paragraph("评级图表：")

    doc.save(output_path)

    with zipfile.ZipFile(output_path, 'a') as z:
        chart_xml = CHART_XML_WITHOUT_JINJA
        z.writestr('word/charts/chart2.xml', chart_xml)

    print(f"已生成: {output_path}")


def create_doc_with_placeholder(output_path: str):
    """创建带占位符的文档。"""
    doc = Document()

    doc.add_heading("研究报告", 0)

    # 添加带占位符的段落
    logo_para = doc.add_paragraph()
    logo_para.add_run("公司 Logo: ")
    logo_run = logo_para.add_run("{{IMAGE}}")
    logo_run.bold = True

    doc.add_paragraph("威胜控股 (3393.HK)")
    doc.add_paragraph("评级：买入")

    doc.save(output_path)
    print(f"已生成: {output_path}")


def create_doc_with_table_placeholder(output_path: str):
    """创建表格中带占位符的文档。"""
    doc = Document()

    doc.add_heading("研究报告", 0)

    # 创建带占位符的表格
    table = doc.add_table(rows=2, cols=2)
    table.style = 'Table Grid'

    table.rows[0].cells[0].text = "公司 Logo"
    table.rows[0].cells[1].text = "{{IMAGE}}"
    table.rows[1].cells[0].text = "评级"
    table.rows[1].cells[1].text = "买入"

    doc.save(output_path)
    print(f"已生成: {output_path}")


def create_source_doc_with_media(output_path: str):
    """创建带媒体文件的源文档（用于测试复制媒体）。"""
    doc = Document()
    doc.add_heading("源文档", 0)
    doc.add_paragraph("这是一个带图片的文档")

    doc.save(output_path)

    # 创建一个简单的图片文件
    # 使用一个最小的有效 PNG 文件（1x1 透明像素）
    png_data = (
        b'\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01'
        b'\x00\x00\x00\x01\x08\x06\x00\x00\x00\x1f\x15\xc4\x89'
        b'\x00\x00\x00\nIDATx\x9cc\x00\x01\x00\x00\x05\x00\x01'
        b'\r\n-\xb4\x00\x00\x00\x00IEND\xaeB`\x82'
    )

    with zipfile.ZipFile(output_path, 'a') as z:
        z.writestr('word/media/image1.png', png_data)

    print(f"已生成: {output_path}")


def create_simple_image():
    """创建一个简单的测试图片文件。"""
    # 这是一个最小的有效 PNG 文件
    png_data = (
        b'\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01'
        b'\x00\x00\x00\x01\x08\x06\x00\x00\x00\x1f\x15\xc4\x89'
        b'\x00\x00\x00\nIDATx\x9cc\x00\x01\x00\x00\x05\x00\x01'
        b'\r\n-\xb4\x00\x00\x00\x00IEND\xaeB`\x82'
    )
    return png_data


if __name__ == "__main__":
    fixtures_dir = Path(__file__).parent
    fixtures_dir.mkdir(exist_ok=True)

    # 生成 image_handler 测试用的 fixture
    create_doc_with_chart_and_jinja(str(fixtures_dir / "doc_with_chart_jinja.docx"))
    create_doc_with_chart_without_jinja(str(fixtures_dir / "doc_with_chart_no_jinja.docx"))
    create_doc_with_placeholder(str(fixtures_dir / "doc_with_placeholder.docx"))
    create_doc_with_table_placeholder(str(fixtures_dir / "doc_with_table_placeholder.docx"))
    create_source_doc_with_media(str(fixtures_dir / "source_with_media.docx"))

    # 生成测试图片
    image_path = fixtures_dir / "test_image.png"
    image_path.write_bytes(create_simple_image())
    print(f"已生成: {image_path}")

    print("\n生成完成！")
