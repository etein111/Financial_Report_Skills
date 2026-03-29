"""模板验证函数。

提供完整性验证和格式验证功能：
- 完整性验证：检查 schema 中所有字段在模板中都有对应占位符
- 格式验证：检查占位符格式（字体大小、粗体、颜色）与原始样本一致
"""

from pathlib import Path
from typing import Dict, Any, List, Tuple, Optional
from dataclasses import dataclass
import re

from docx import Document
from docx.shared import Pt, RGBColor

import common


# ==================== 数据结构 ====================

@dataclass
class FormatInfo:
    """占位符格式信息。"""
    font_size: Optional[float] = None  # 字号（磅）
    bold: Optional[bool] = None  # 是否粗体
    color_rgb: Optional[str] = None  # 颜色 RGB（如 "000000"）
    font_name: Optional[str] = None  # 字体名称


@dataclass
class ValidationResult:
    """验证结果。"""
    passed: bool
    message: str
    details: List[str] = None

    def __post_init__(self):
        if self.details is None:
            self.details = []


# ==================== 占位符提取 ====================

def extract_placeholders_from_doc(doc_path: str) -> Dict[str, FormatInfo]:
    """从模板文档中提取所有占位符及其格式信息。

    Args:
        doc_path: 模板文档路径

    Returns:
        占位符名称 -> 格式信息 的映射
    """
    doc = Document(doc_path)
    placeholders = {}

    # 匹配 {{ field_name }} 或 [Table_XXX] 格式
    pattern_jinja = re.compile(r'\{\{\s*(\w+)\s*\}\}')
    pattern_table = re.compile(r'\[Table_(\w+)\]')

    def _extract_format(run) -> FormatInfo:
        """从 run 提取格式信息。"""
        font = run.font
        color_rgb = None
        if font.color and font.color.rgb:
            color_rgb = str(font.color.rgb)

        return FormatInfo(
            font_size=font.size.pt if font.size else None,
            bold=font.bold,
            color_rgb=color_rgb,
            font_name=font.name
        )

    def _process_paragraph(para):
        """处理段落，提取占位符。"""
        for run in para.runs:
            text = run.text
            # 检查 Jinja 占位符
            for match in pattern_jinja.finditer(text):
                field_name = match.group(1)
                if field_name not in placeholders:
                    placeholders[field_name] = _extract_format(run)

            # 检查 Table 占位符
            for match in pattern_table.finditer(text):
                field_name = f"Table_{match.group(1)}"
                if field_name not in placeholders:
                    placeholders[field_name] = _extract_format(run)

    def _process_table(table):
        """递归处理表格。"""
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    _process_paragraph(para)
                # 递归处理嵌套表格
                for nested_table in cell.tables:
                    _process_table(nested_table)

    # 处理文档段落
    for para in doc.paragraphs:
        _process_paragraph(para)

    # 处理所有表格
    for table in doc.tables:
        _process_table(table)

    # 处理页眉页脚
    for section in doc.sections:
        # 页眉
        for header in [section.header, section.first_page_header, section.even_page_header]:
            if header and header.is_linked_to_previous is False:
                for para in header.paragraphs:
                    _process_paragraph(para)
                for table in header.tables:
                    _process_table(table)

        # 页脚
        for footer in [section.footer, section.first_page_footer, section.even_page_footer]:
            if footer and footer.is_linked_to_previous is False:
                for para in footer.paragraphs:
                    _process_paragraph(para)
                for table in footer.tables:
                    _process_table(table)

    return placeholders


def get_expected_fields_from_schema(schema: Dict[str, Any]) -> List[str]:
    """从 schema 获取预期的字段列表。

    Args:
        schema: schema.yaml 内容

    Returns:
        预期字段名列表
    """
    fields = []

    for field_name, field_config in schema.items():
        if isinstance(field_config, dict):
            field_type = field_config.get('type')

            if field_type == 'fragment':
                # fragment 类型：abstract, body, appendix
                fields.append(field_name)
                # appendix 还有 Table 占位符
                if field_name == 'appendix':
                    fields.append('Table_APPENDIX')
            elif field_name == 'sections':
                # sections 是复合字段，检查子字段
                # 不添加 'sections' 本身，只添加 'sections_*' 子字段
                format_str = field_config.get('format', '')
                try:
                    import ast
                    format_dict = ast.literal_eval(format_str)
                    if isinstance(format_dict, dict):
                        for sub_key in format_dict.keys():
                            fields.append(f'sections_{sub_key}')
                except (ValueError, SyntaxError):
                    # 如果解析失败，添加 sections 本身
                    fields.append(field_name)
            elif field_name == 'analysts':
                # analysts 是列表字段，检查展平后的子字段
                # 不添加 'analysts' 本身，只添加 'analyst_*' 子字段
                format_str = field_config.get('format', '')
                try:
                    import ast
                    analysts_list = ast.literal_eval(format_str)
                    if isinstance(analysts_list, list) and analysts_list:
                        first = analysts_list[0]
                        if isinstance(first, dict):
                            for key in first.keys():
                                fields.append(f'analyst_{key}')
                except (ValueError, SyntaxError):
                    # 如果解析失败，添加 analysts 本身
                    fields.append(field_name)
            else:
                # 普通字段
                fields.append(field_name)
        else:
            # 简单字段
            fields.append(field_name)

    return fields


# ==================== 完整性验证 ====================

def validate_completeness(
    template_path: str,
    schema_path: str,
    skip_fields: List[str] = None
) -> ValidationResult:
    """验证模板完整性：schema 中的字段是否都有对应占位符。

    Args:
        template_path: 模板文档路径
        schema_path: schema.yaml 路径
        skip_fields: 跳过验证的字段列表（用于特殊位置的字段）

    Returns:
        验证结果
    """
    skip_fields = skip_fields or []

    try:
        # 加载 schema
        schema = common.load_yaml(schema_path)
        if not schema:
            return ValidationResult(
                passed=False,
                message="schema.yaml 为空或不存在"
            )

        # 获取预期字段
        expected_fields = get_expected_fields_from_schema(schema)

        # 移除跳过的字段
        expected_fields = [f for f in expected_fields if f not in skip_fields]

        # 提取模板中的占位符
        placeholders = extract_placeholders_from_doc(template_path)
        found_fields = set(placeholders.keys())

        # 检查缺失的字段
        missing = []
        for field in expected_fields:
            if field not in found_fields:
                missing.append(field)

        # 检查多余的占位符（可选警告）
        extra = found_fields - set(expected_fields) - set(skip_fields)

        if missing:
            return ValidationResult(
                passed=False,
                message=f"缺少 {len(missing)} 个占位符",
                details=[f"缺失: {f}" for f in missing]
            )

        details = [f"✓ 找到 {len(found_fields)} 个占位符，schema 定义 {len(expected_fields)} 个字段"]
        if skip_fields:
            details.append(f"跳过验证: {', '.join(skip_fields)}")
        if extra:
            details.append(f"额外占位符（可忽略）: {', '.join(extra)}")

        return ValidationResult(
            passed=True,
            message="完整性验证通过",
            details=details
        )

    except Exception as e:
        return ValidationResult(
            passed=False,
            message=f"验证出错: {str(e)}"
        )


# ==================== 格式验证 ====================

def extract_format_from_sample(
    sample_path: str,
    field_values: Dict[str, str]
) -> Dict[str, FormatInfo]:
    """从原始样本中提取字段值的格式信息。

    通过在样本中查找字段值，获取其格式信息。

    Args:
        sample_path: 原始样本文档路径
        field_values: 字段名 -> 字段值 的映射

    Returns:
        字段名 -> 格式信息 的映射
    """
    doc = Document(sample_path)
    format_info = {}

    def _extract_format(run) -> FormatInfo:
        """从 run 提取格式信息。"""
        font = run.font
        color_rgb = None
        if font.color and font.color.rgb:
            color_rgb = str(font.color.rgb)

        return FormatInfo(
            font_size=font.size.pt if font.size else None,
            bold=font.bold,
            color_rgb=color_rgb,
            font_name=font.name
        )

    def _search_in_paragraph(para, field_name: str, field_value: str) -> bool:
        """在段落中搜索字段值。"""
        # 检查段落文本是否包含字段值
        para_text = para.text
        if field_value and len(field_value) > 10 and field_value[:20] in para_text:
            # 找到了，提取第一个 run 的格式
            if para.runs:
                format_info[field_name] = _extract_format(para.runs[0])
                return True
        return False

    def _search_in_table(table, field_name: str, field_value: str) -> bool:
        """在表格中搜索字段值。"""
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    if _search_in_paragraph(para, field_name, field_value):
                        return True
                for nested_table in cell.tables:
                    if _search_in_table(nested_table, field_name, field_value):
                        return True
        return False

    # 搜索每个字段值
    for field_name, field_value in field_values.items():
        if not field_value or len(field_value) < 10:
            continue

        # 先在段落中搜索
        for para in doc.paragraphs:
            if _search_in_paragraph(para, field_name, field_value):
                break

        # 如果没找到，在表格中搜索
        if field_name not in format_info:
            for table in doc.tables:
                if _search_in_table(table, field_name, field_value):
                    break

    return format_info


def validate_format(
    template_path: str,
    sample_path: str,
    elements_path: str,
    skip_fields: List[str] = None
) -> ValidationResult:
    """验证模板格式：占位符格式是否与原始样本一致。

    Args:
        template_path: 模板文档路径
        sample_path: 原始样本文档路径
        elements_path: elements.yaml 路径
        skip_fields: 跳过格式验证的字段列表

    Returns:
        验证结果
    """
    skip_fields = skip_fields or []

    try:
        # 加载 elements
        elements = common.load_yaml(elements_path)
        if not elements:
            return ValidationResult(
                passed=False,
                message="elements.yaml 为空或不存在"
            )

        # 展平 elements
        flat_elements = common.flatten_elements(elements)

        # 从样本提取格式信息
        sample_formats = extract_format_from_sample(sample_path, flat_elements)

        # 从模板提取占位符格式
        template_formats = extract_placeholders_from_doc(template_path)

        # 比较格式
        mismatches = []
        checked = 0
        skipped = 0

        for field_name, sample_fmt in sample_formats.items():
            # 跳过指定字段
            if field_name in skip_fields:
                skipped += 1
                continue

            template_fmt = template_formats.get(field_name)
            if not template_fmt:
                continue

            checked += 1
            issues = []

            # 比较字体大小
            if sample_fmt.font_size and template_fmt.font_size:
                if abs(sample_fmt.font_size - template_fmt.font_size) > 0.5:
                    issues.append(
                        f"字号: 模板 {template_fmt.font_size}pt vs 样本 {sample_fmt.font_size}pt"
                    )

            # 比较粗体
            if sample_fmt.bold is not None and template_fmt.bold is not None:
                if sample_fmt.bold != template_fmt.bold:
                    issues.append(
                        f"粗体: 模板 {template_fmt.bold} vs 样本 {sample_fmt.bold}"
                    )

            # 比较颜色
            if sample_fmt.color_rgb and template_fmt.color_rgb:
                if sample_fmt.color_rgb != template_fmt.color_rgb:
                    issues.append(
                        f"颜色: 模板 #{template_fmt.color_rgb} vs 样本 #{sample_fmt.color_rgb}"
                    )

            if issues:
                mismatches.append(f"{field_name}: {'; '.join(issues)}")

        if mismatches:
            return ValidationResult(
                passed=False,
                message=f"发现 {len(mismatches)} 个格式不一致",
                details=mismatches
            )

        details = [f"✓ 已验证 {checked} 个字段的格式"]
        if skipped:
            details.append(f"跳过 {skipped} 个字段")

        return ValidationResult(
            passed=True,
            message=f"格式验证通过（检查了 {checked} 个字段）",
            details=details
        )

    except Exception as e:
        return ValidationResult(
            passed=False,
            message=f"格式验证出错: {str(e)}"
        )


# ==================== 综合验证 ====================

# 默认跳过验证的字段（通常在首页表格等特殊位置）
DEFAULT_SKIP_FIELDS = ['ticker', 'title_ch', 'title_en']


def validate_template(
    template_path: str,
    schema_path: str,
    sample_path: str = None,
    elements_path: str = None,
    skip_fields: List[str] = None,
    skip_format_fields: List[str] = None
) -> Tuple[bool, str, Dict[str, ValidationResult]]:
    """综合验证模板。

    Args:
        template_path: 模板文档路径
        schema_path: schema.yaml 路径
        sample_path: 原始样本路径（可选，用于格式验证）
        elements_path: elements.yaml 路径（可选，用于格式验证）
        skip_fields: 跳过完整性验证的字段
        skip_format_fields: 跳过格式验证的字段

    Returns:
        (是否全部通过, 总结消息, 各项验证结果)
    """
    # 使用默认跳过字段
    if skip_fields is None:
        skip_fields = DEFAULT_SKIP_FIELDS.copy()

    skip_format_fields = skip_format_fields or []

    results = {}

    # 1. 完整性验证
    completeness_result = validate_completeness(
        template_path, schema_path, skip_fields=skip_fields
    )
    results['completeness'] = completeness_result

    # 2. 格式验证（如果提供了样本和 elements）
    if sample_path and elements_path:
        format_result = validate_format(
            template_path, sample_path, elements_path,
            skip_fields=skip_format_fields
        )
        results['format'] = format_result
    else:
        results['format'] = ValidationResult(
            passed=True,
            message="跳过格式验证（未提供样本或 elements）"
        )

    # 汇总结果
    all_passed = all(r.passed for r in results.values())

    if all_passed:
        summary = "✓ 模板验证通过"
    else:
        failed = [name for name, r in results.items() if not r.passed]
        summary = f"✗ 模板验证失败: {', '.join(failed)}"

    return all_passed, summary, results


def print_validation_report(results: Dict[str, ValidationResult]) -> None:
    """打印验证报告。

    Args:
        results: 验证结果字典
    """
    print("\n" + "=" * 50)
    print("模板验证报告")
    print("=" * 50)

    for name, result in results.items():
        status = "✓" if result.passed else "✗"
        print(f"\n{status} {name}: {result.message}")
        if result.details:
            for detail in result.details:
                print(f"  - {detail}")

    print("\n" + "=" * 50)
