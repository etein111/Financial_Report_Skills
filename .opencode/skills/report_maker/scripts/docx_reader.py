"""docx_reader.py - Word 文档读取模块。

提供递归读取嵌套表格的功能，解决 python-docx 的 cell.text
不包含嵌套表格内容的问题。

遵循函数式编程原则：纯函数、无副作用、I/O 分离。
"""

from typing import List, Dict, Any, Optional
from pathlib import Path


def read_docx_tables(docx_path: str) -> List[Dict[str, Any]]:
    """读取 Word 文档中的所有表格（包括嵌套表格）。

    Args:
        docx_path: Word 文档路径

    Returns:
        表格列表，每个表格包含结构化的数据

    Raises:
        FileNotFoundError: 文件不存在
        Exception: 文档读取错误
    """
    from docx import Document

    doc_path = Path(docx_path)
    if not doc_path.exists():
        raise FileNotFoundError(f"文件不存在: {docx_path}")

    doc = Document(docx_path)
    tables = []

    for table_idx, table in enumerate(doc.tables):
        table_data = _extract_table_data(table, table_idx)
        tables.append(table_data)

    return tables


def _extract_table_data(table, table_idx: int) -> Dict[str, Any]:
    """提取单个表格的数据（包括嵌套表格）。

    Args:
        table: python-docx Table 对象
        table_idx: 表格索引

    Returns:
        表格数据字典
    """
    rows_data = []

    for row_idx, row in enumerate(table.rows):
        row_data = _extract_row_data(row, row_idx)
        rows_data.append(row_data)

    return {
        'index': table_idx,
        'type': 'table',
        'rows': rows_data
    }


def _extract_row_data(row, row_idx: int) -> Dict[str, Any]:
    """提取行数据（包括嵌套表格）。

    Args:
        row: python-docx Row 对象
        row_idx: 行索引

    Returns:
        行数据字典
    """
    cells_data = []

    for cell_idx, cell in enumerate(row.cells):
        cell_data = _extract_cell_data(cell, cell_idx)
        cells_data.append(cell_data)

    return {
        'index': row_idx,
        'type': 'row',
        'cells': cells_data
    }


def _extract_cell_data(cell, cell_idx: int) -> Dict[str, Any]:
    """提取单元格数据（包括嵌套表格）。

    Args:
        cell: python-docx Cell 对象
        cell_idx: 单元格索引

    Returns:
        单元格数据字典
    """
    # 提取段落文本
    paragraphs_text = []
    for para in cell.paragraphs:
        text = para.text.strip()
        if text:
            paragraphs_text.append(text)

    # 提取嵌套表格
    nested_tables = []
    for nested_idx, nested_table in enumerate(cell.tables):
        nested_data = _extract_table_data(nested_table, nested_idx)
        nested_tables.append(nested_data)

    return {
        'index': cell_idx,
        'type': 'cell',
        'text': '\n'.join(paragraphs_text),
        'paragraphs': paragraphs_text,
        'nested_tables': nested_tables
    }


def flatten_table_text(tables: List[Dict[str, Any]]) -> List[str]:
    """扁平化表格数据为文本列表。

    用于快速搜索和匹配字段。

    Args:
        tables: 表格数据列表

    Returns:
        文本行列表
    """
    lines = []

    def _flatten_node(node: Dict[str, Any], prefix: str = "") -> None:
        """递归扁平化节点。"""
        node_type = node.get('type', '')

        if node_type == 'cell':
            # 添加单元格文本
            text = node.get('text', '').strip()
            if text:
                lines.append(text)

            # 递归处理嵌套表格
            for nested_table in node.get('nested_tables', []):
                _flatten_node(nested_table, f"{prefix}  ")

        elif node_type == 'row':
            for cell in node.get('cells', []):
                _flatten_node(cell, prefix)

        elif node_type == 'table':
            for row in node.get('rows', []):
                _flatten_node(row, prefix)

    # 处理所有表格
    for table in tables:
        _flatten_node(table)

    return lines


def search_in_tables(tables: List[Dict[str, Any]], keyword: str) -> List[Dict[str, Any]]:
    """在表格中搜索关键词。

    Args:
        tables: 表格数据列表
        keyword: 搜索关键词

    Returns:
        匹配的单元格列表
    """
    matches = []
    keyword_lower = keyword.lower()

    def _search_node(node: Dict[str, Any]) -> None:
        """递归搜索节点。"""
        node_type = node.get('type', '')

        if node_type == 'cell':
            text = node.get('text', '').strip()
            if keyword_lower in text.lower():
                matches.append(node)

            # 递归搜索嵌套表格
            for nested_table in node.get('nested_tables', []):
                _search_node(nested_table)

        elif node_type == 'row':
            for cell in node.get('cells', []):
                _search_node(cell)

        elif node_type == 'table':
            for row in node.get('rows', []):
                _search_node(row)

    # 搜索所有表格
    for table in tables:
        _search_node(table)

    return matches


def extract_nested_tables_only(tables: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    """仅提取嵌套表格的数据。

    用于快速定位首页表格中的嵌套表格（通常包含关键字段）。

    Args:
        tables: 表格数据列表

    Returns:
        嵌套表格列表
    """
    nested_tables = []

    def _collect_nested(node: Dict[str, Any]) -> None:
        """递归收集嵌套表格。"""
        node_type = node.get('type', '')

        if node_type == 'cell':
            for nested_table in node.get('nested_tables', []):
                nested_tables.append(nested_table)
                # 递归处理更深层的嵌套
                _collect_nested(nested_table)

        elif node_type == 'row':
            for cell in node.get('cells', []):
                _collect_nested(cell)

        elif node_type == 'table':
            for row in node.get('rows', []):
                _collect_nested(row)

    # 收集所有嵌套表格
    for table in tables:
        _collect_nested(table)

    return nested_tables
