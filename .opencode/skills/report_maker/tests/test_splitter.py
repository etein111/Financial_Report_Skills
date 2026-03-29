"""测试 splitter.py 模块的函数。"""
import pytest
from pathlib import Path
import tempfile
import shutil

import sys
sys.path.insert(0, str(Path(__file__).parent.parent / "scripts"))

from splitter import (
    extract_abstract, extract_body, get_split_instructions,
    locate_fragment_by_content, extract_fragment_by_table_range,
    extract_fragment_by_paragraph_range, extract_fragment_by_keywords,
    read_document_structure
)


class TestExtractAbstract:
    """测试 extract_abstract 函数。"""

    def test_returns_path(self):
        """测试：返回 abstract 文件路径。"""
        result = extract_abstract("draft.docx", "首页右侧")

        assert result == Path("abstract.docx")


class TestExtractBody:
    """测试 extract_body 函数。"""

    def test_returns_path(self):
        """测试：返回 body 文件路径。"""
        result = extract_body("draft.docx", "首页之后")

        assert result == Path("body.docx")


class TestGetSplitInstructions:
    """测试 get_split_instructions 函数。"""

    def test_abstract_instructions(self):
        """测试：生成摘要拆分指令。"""
        schema = {
            "abstract": {
                "type": "fragment",
                "scope": "首页右侧",
                "hints": ["通常包含投资亮点", "长度约 200-500 字"]
            }
        }

        instructions = get_split_instructions(schema)

        assert "摘要位置：首页右侧" in instructions
        assert "摘要特征" in instructions

    def test_body_instructions(self):
        """测试：生成正文拆分指令。"""
        schema = {
            "body": {
                "type": "fragment",
                "scope": "从第三页开始",
                "hints": ["包含各章节详细内容"]
            }
        }

        instructions = get_split_instructions(schema)

        assert "正文位置：从第三页开始" in instructions
        assert "正文特征" in instructions

    def test_both_instructions(self):
        """测试：同时生成摘要和正文指令。"""
        schema = {
            "abstract": {
                "type": "fragment",
                "scope": "首页右侧",
                "hints": ["投资亮点"]
            },
            "body": {
                "type": "fragment",
                "scope": "第三页开始",
                "hints": ["详细内容"]
            }
        }

        instructions = get_split_instructions(schema)

        assert "摘要位置：首页右侧" in instructions
        assert "正文位置：第三页开始" in instructions

    def test_no_fragment_fields(self):
        """测试：没有 fragment 字段时返回空字符串。"""
        schema = {
            "ticker": {},
            "company": {}
        }

        instructions = get_split_instructions(schema)

        assert instructions == ""


class TestLocateFragmentByContent:
    """测试 locate_fragment_by_content 函数。"""

    @pytest.fixture
    def sample_doc(self, tmp_path):
        """创建示例文档。"""
        try:
            from docx import Document

            doc = Document()

            # 添加一些段落
            doc.add_paragraph("这是标题")
            doc.add_paragraph("投资亮点")
            doc.add_paragraph("这是正文内容")
            doc.add_paragraph("风险提示")

            # 添加表格
            table = doc.add_table(rows=2, cols=2)
            table.rows[0].cells[0].text = "财务数据"
            table.rows[0].cells[1].text = "2024"
            table.rows[1].cells[0].text = "收入"
            table.rows[1].cells[1].text = "1000万"

            doc_path = tmp_path / "sample.docx"
            doc.save(str(doc_path))

            return str(doc_path)
        except ImportError:
            pytest.skip("python-docx not installed")

    def test_locate_by_keywords_in_paragraphs(self, sample_doc):
        """测试：在段落中定位关键词。"""
        result = locate_fragment_by_content(sample_doc, ["投资亮点"])

        assert "paragraphs" in result
        assert len(result["paragraphs"]) == 1
        assert result["paragraphs"][0]["keyword"] == "投资亮点"
        assert "投资亮点" in result["paragraphs"][0]["text"]

    def test_locate_by_keywords_in_tables(self, sample_doc):
        """测试：在表格中定位关键词。"""
        result = locate_fragment_by_content(sample_doc, ["财务数据"])

        assert "tables" in result
        assert len(result["tables"]) == 1
        assert result["tables"][0]["matched_cells"][0]["keyword"] == "财务数据"

    def test_locate_multiple_keywords(self, sample_doc):
        """测试：使用多个关键词搜索。"""
        result = locate_fragment_by_content(sample_doc, ["投资亮点", "风险提示"])

        assert result["total_matches"] >= 2
        assert len(result["paragraphs"]) >= 2

    def test_no_matches(self, sample_doc):
        """测试：没有匹配的情况。"""
        result = locate_fragment_by_content(sample_doc, ["不存在的内容"])

        assert result["total_matches"] == 0
        assert len(result["paragraphs"]) == 0
        assert len(result["tables"]) == 0


class TestExtractFragmentByTableRange:
    """测试 extract_fragment_by_table_range 函数。"""

    @pytest.fixture
    def sample_doc(self, tmp_path):
        """创建包含多个表格的示例文档。"""
        try:
            from docx import Document

            doc = Document()
            doc.add_paragraph("文档开始")

            # 表格 1
            table1 = doc.add_table(rows=2, cols=2)
            table1.rows[0].cells[0].text = "表格1"
            table1.rows[0].cells[1].text = "数据1"

            # 表格 2
            table2 = doc.add_table(rows=2, cols=2)
            table2.rows[0].cells[0].text = "表格2"
            table2.rows[0].cells[1].text = "数据2"

            # 表格 3
            table3 = doc.add_table(rows=2, cols=2)
            table3.rows[0].cells[0].text = "表格3"
            table3.rows[0].cells[1].text = "数据3"

            doc.add_paragraph("文档结束")

            doc_path = tmp_path / "sample.docx"
            doc.save(str(doc_path))

            return str(doc_path)
        except ImportError:
            pytest.skip("python-docx not installed")

    def test_extract_single_table(self, sample_doc, tmp_path):
        """测试：提取单个表格。"""
        output_path = tmp_path / "output.docx"

        result = extract_fragment_by_table_range(sample_doc, 0, 0, str(output_path))

        assert result is True
        assert Path(output_path).exists()

    def test_extract_multiple_tables(self, sample_doc, tmp_path):
        """测试：提取多个表格范围。"""
        output_path = tmp_path / "output.docx"

        result = extract_fragment_by_table_range(sample_doc, 0, 1, str(output_path))

        assert result is True
        assert Path(output_path).exists()

    def test_invalid_range(self, sample_doc, tmp_path):
        """测试：无效的表格范围。"""
        output_path = tmp_path / "output.docx"

        result = extract_fragment_by_table_range(sample_doc, 10, 20, str(output_path))

        assert result is False


class TestExtractFragmentByParagraphRange:
    """测试 extract_fragment_by_paragraph_range 函数。"""

    @pytest.fixture
    def sample_doc(self, tmp_path):
        """创建示例文档。"""
        try:
            from docx import Document

            doc = Document()

            for i in range(10):
                doc.add_paragraph(f"段落 {i}")

            doc_path = tmp_path / "sample.docx"
            doc.save(str(doc_path))

            return str(doc_path)
        except ImportError:
            pytest.skip("python-docx not installed")

    def test_extract_paragraph_range(self, sample_doc, tmp_path):
        """测试：提取指定段落范围。"""
        output_path = tmp_path / "output.docx"

        result = extract_fragment_by_paragraph_range(sample_doc, 2, 5, str(output_path))

        assert result is True
        assert Path(output_path).exists()

    def test_extract_from_start(self, sample_doc, tmp_path):
        """测试：从开头提取。"""
        output_path = tmp_path / "output.docx"

        result = extract_fragment_by_paragraph_range(sample_doc, 0, 3, str(output_path))

        assert result is True
        assert Path(output_path).exists()

    def test_extract_to_end(self, sample_doc, tmp_path):
        """测试：提取到末尾。"""
        output_path = tmp_path / "output.docx"

        result = extract_fragment_by_paragraph_range(sample_doc, 5, 100, str(output_path))

        assert result is True
        assert Path(output_path).exists()


class TestExtractFragmentByKeywords:
    """测试 extract_fragment_by_keywords 函数。"""

    @pytest.fixture
    def sample_doc(self, tmp_path):
        """创建示例文档。"""
        try:
            from docx import Document

            doc = Document()

            doc.add_paragraph("开始标记")
            doc.add_paragraph("这是要提取的内容")
            doc.add_paragraph("更多内容")
            doc.add_paragraph("结束标记")

            doc_path = tmp_path / "sample.docx"
            doc.save(str(doc_path))

            return str(doc_path)
        except ImportError:
            pytest.skip("python-docx not installed")

    def test_extract_by_keywords(self, sample_doc, tmp_path):
        """测试：根据关键词提取片段。"""
        output_path = tmp_path / "output.docx"

        result = extract_fragment_by_keywords(
            sample_doc,
            ["开始标记"],
            ["结束标记"],
            str(output_path)
        )

        assert result is True
        assert Path(output_path).exists()

    def test_keywords_not_found(self, sample_doc, tmp_path):
        """测试：关键词未找到。"""
        output_path = tmp_path / "output.docx"

        result = extract_fragment_by_keywords(
            sample_doc,
            ["不存在的内容"],
            ["结束标记"],
            str(output_path)
        )

        assert result is False


class TestReadDocumentStructure:
    """测试 read_document_structure 函数。"""

    @pytest.fixture
    def sample_doc(self, tmp_path):
        """创建示例文档。"""
        try:
            from docx import Document

            doc = Document()

            for i in range(5):
                doc.add_paragraph(f"段落 {i}")

            table = doc.add_table(rows=2, cols=2)
            table.rows[0].cells[0].text = "表头"

            doc_path = tmp_path / "sample.docx"
            doc.save(str(doc_path))

            return str(doc_path)
        except ImportError:
            pytest.skip("python-docx not installed")

    def test_read_structure(self, sample_doc):
        """测试：读取文档结构。"""
        result = read_document_structure(sample_doc)

        assert "paragraphs" in result
        assert "tables" in result
        assert result["paragraphs"] == 5
        assert result["tables"] == 1

    def test_paragraph_preview(self, sample_doc):
        """测试：段落预览。"""
        result = read_document_structure(sample_doc)

        assert "paragraph_preview" in result
        assert len(result["paragraph_preview"]) == 5
        assert "text" in result["paragraph_preview"][0]
        assert "index" in result["paragraph_preview"][0]

    def test_table_preview(self, sample_doc):
        """测试：表格预览。"""
        result = read_document_structure(sample_doc)

        assert "table_preview" in result
        assert len(result["table_preview"]) == 1
        assert "rows" in result["table_preview"][0]
        assert "cols" in result["table_preview"][0]
