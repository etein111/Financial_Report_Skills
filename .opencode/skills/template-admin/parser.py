"""Word 文档解析器"""
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from typing import Dict, List, Any, Tuple
import re


class WordParser:
    """解析 Word 文档结构"""
    
    def __init__(self, file_path: str):
        self.file_path = file_path
        self.doc = Document(file_path)
    
    def parse_structure(self) -> Dict[str, Any]:
        """解析完整结构"""
        return {
            'tables': self.extract_tables(),
            'paragraphs': self.extract_paragraphs(),
            'styles': self.extract_styles(),
            'metadata': self.extract_metadata()
        }
    
    def extract_tables(self) -> List[Dict]:
        """提取所有表格"""
        tables = []
        for idx, table in enumerate(self.doc.tables):
            rows_data = []
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                rows_data.append(cells)
            
            tables.append({
                'index': idx,
                'rows': rows_data,
                'row_count': len(table.rows),
                'col_count': len(table.columns) if table.rows else 0
            })
        return tables
    
    def extract_paragraphs(self) -> List[Dict]:
        """提取所有段落"""
        paragraphs = []
        for idx, para in enumerate(self.doc.paragraphs):
            text = para.text.strip()
            if text:
                paragraphs.append({
                    'index': idx,
                    'text': text,
                    'style': para.style.name if para.style else 'Normal',
                    'is_heading': para.style.name.startswith('Heading')
                })
        return paragraphs
    
    def extract_styles(self) -> Dict[str, Any]:
        """提取样式信息"""
        styles = {}
        for style in self.doc.styles:
            if style.type == 1:  # Paragraph style
                styles[style.name] = {
                    'type': 'paragraph',
                    'base_style': style.base_style.name if style.base_style else None
                }
        return styles
    
    def extract_metadata(self) -> Dict:
        """提取文档元数据"""
        core = self.doc.core_properties
        return {
            'author': core.author,
            'title': core.title,
            'created': core.created.isoformat() if core.created else None,
            'modified': core.modified.isoformat() if core.modified else None
        }
    
    def find_dynamic_regions(self) -> List[Dict]:
        """查找动态内容区域（可能被替换的内容）"""
        regions = []
        
        # 查找可能是占位符的文本模式
        pattern = r'\{[^}]+\}|\[[^\]]+\]|<[^>]+>'
        
        for para in self.extract_paragraphs():
            matches = re.findall(pattern, para['text'])
            if matches:
                regions.append({
                    'type': 'paragraph',
                    'index': para['index'],
                    'matches': matches
                })
        
        # 查找表格中的动态内容
        for table in self.extract_tables():
            for row_idx, row in enumerate(table['rows']):
                for col_idx, cell in enumerate(row):
                    matches = re.findall(pattern, cell)
                    if matches:
                        regions.append({
                            'type': 'table',
                            'table_index': table['index'],
                            'row': row_idx,
                            'col': col_idx,
                            'matches': matches
                        })
        
        return regions
