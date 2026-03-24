"""Template.docx 生成器。"""

from pathlib import Path
from docx import Document
from typing import Dict

class TemplateBuilder:
    """从 Word 样本生成 template.docx。"""

    def build(self, sample_path: Path, placeholders: Dict[str, str], output_path: Path):
        """构建模板文档。

        Args:
            sample_path: 样本文档路径
            placeholders: 占位符映射 {位置: jinja2_var}
            output_path: 输出路径
        """
        doc = Document(sample_path)

        # 替换占位符
        # 简化版本：替换表格中的文本
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for key, var in placeholders.items():
                        if key in cell.text:
                            # 替换为 Jinja2 占位符
                            cell.text = cell.text.replace(key, var)

        # 保存
        output_path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(output_path)
