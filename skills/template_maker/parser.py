"""Word 文档结构分析器（提示词包装）。"""

from pathlib import Path
from docx import Document
from typing import Dict, Any

class WordStructureAnalyzer:
    """分析 Word 文档结构。"""

    def analyze(self, doc_path: Path) -> Dict[str, Any]:
        """分析 Word 文档结构。

        Args:
            doc_path: Word 文档路径

        Returns:
            包含以下键的字典：
            - tables: 表格列表
            - paragraphs: 段落数量
            - key_areas: 关键区域（页眉、页脚等）
        """
        doc = Document(doc_path)

        # 使用 Claude Code 的 LLM 能力分析结构
        # 这里先返回基本信息，后续会用提示词增强
        structure = {
            "tables": self._extract_tables(doc),
            "paragraphs": len(doc.paragraphs),
            "key_areas": {
                "headers": len(doc.sections),
                "footers": len(doc.sections),
            }
        }

        return structure

    def _extract_tables(self, doc: Document) -> list:
        """提取表格信息。"""
        tables = []
        for i, table in enumerate(doc.tables):
            tables.append({
                "index": i,
                "rows": len(table.rows),
                "cols": len(table.columns),
            })
        return tables
