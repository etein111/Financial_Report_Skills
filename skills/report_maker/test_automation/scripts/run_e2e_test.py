#!/usr/bin/env python3
"""
run_e2e_test.py - Report Maker 端到端自动化测试

用法：
    python run_e2e_test.py [--skip-extract] [--skip-render]

选项：
    --skip-extract  跳过第一阶段提取（复用已有输出）
    --skip-render   跳过第二阶段渲染（仅验证提取结果）
"""
import sys
import os
import re
import shutil
import yaml
import importlib.util
from pathlib import Path
from docx import Document
from docxtpl import DocxTemplate
from docxcompose.composer import Composer
from copy import deepcopy

# === Paths ===
DRAFT_PATH = 'inbox/samples/company/0402威胜控股（3393.HK）首次覆盖：配网投资与AI算力双轮驱动，能源计量龙头再启航.docx'
OUTPUT_DIR = 'reports/test_output/'
TEMPLATE_PATH = 'templates/company/zh/template.docx'

# === Expected values (for regression check) ===
# Note: body_paras and eng_summary_paras count ALL paragraphs (incl. empty)
# body_doc.paragraphs includes empty paragraphs from section breaks
EXPECTED = {
    'abstract_chars': (700, 900),
    'abstract_paras': (6, 10),
    'body_all_paras': (260, 285),     # all paragraphs (incl empty)
    'body_tables': (24, 28),
    'eng_summary_all_paras': (17, 21),  # all paragraphs (incl empty)
    'final_paragraphs': (290, 310),
    'final_tables': (26, 30),
}

# === Helper functions ===

def check(name, condition, expected, actual):
    status = 'PASS' if condition else 'FAIL'
    print(f'  [{status}] {name}: expected={expected}, actual={actual}')
    return condition

def load_splitter():
    spec = importlib.util.spec_from_file_location('splitter',
        'skills/report_maker/scripts/splitter.py')
    splitter = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(splitter)
    return splitter

def in_range(val, rng):
    return rng[0] <= val <= rng[1]

# === Phase 1: Extract ===

def phase1_extract():
    print('\n=== Phase 1: Extraction ===')
    Path(OUTPUT_DIR).mkdir(parents=True, exist_ok=True)
    shutil.copy(DRAFT_PATH, f'{OUTPUT_DIR}/draft.docx')

    splitter = load_splitter()

    # Abstract
    splitter.extract_abstract(DRAFT_PATH, '首页右侧表格第8行', f'{OUTPUT_DIR}/abstract.docx')
    abs_doc = Document(f'{OUTPUT_DIR}/abstract.docx')
    abs_chars = sum(len(p.text) for p in abs_doc.paragraphs)
    abs_paras = len([p for p in abs_doc.paragraphs if p.text.strip()])

    check('Abstract chars', in_range(abs_chars, EXPECTED['abstract_chars']), EXPECTED['abstract_chars'], abs_chars)
    check('Abstract paragraphs', in_range(abs_paras, EXPECTED['abstract_paras']), EXPECTED['abstract_paras'], abs_paras)

    # Body
    splitter.extract_body(DRAFT_PATH, 'P9到P273',
        f'{OUTPUT_DIR}/body.docx',
        start_marker='1 威胜控股：能源计量与能效管理专家',
        end_marker='国家电网及南方电网集采招标量价不及预期；海外市场业务增长速度慢于预期；海外生产成本上升')
    body_doc = Document(f'{OUTPUT_DIR}/body.docx')
    body_all_paras = len(body_doc.paragraphs)
    body_tables = len(body_doc.tables)

    check('Body all paragraphs', in_range(body_all_paras, EXPECTED['body_all_paras']), EXPECTED['body_all_paras'], body_all_paras)
    check('Body tables', in_range(body_tables, EXPECTED['body_tables']), EXPECTED['body_tables'], body_tables)

    # English summary
    splitter.extract_english_summary(DRAFT_PATH, 'P275到P291',
        f'{OUTPUT_DIR}/english_summary.docx',
        start_marker='Initial Coverage on Wasion Holdings')
    eng_doc = Document(f'{OUTPUT_DIR}/english_summary.docx')
    eng_all_paras = len(eng_doc.paragraphs)

    check('English summary all paragraphs', in_range(eng_all_paras, EXPECTED['eng_summary_all_paras']), EXPECTED['eng_summary_all_paras'], eng_all_paras)

    # elements.yaml
    elements = {
        'report_date': '2 Apr 2026',
        'company_cn': '威胜控股',
        'ticker': '3393.HK',
        'title': '首次覆盖：配网投资与AI算力双轮驱动，能源计量龙头再启航',
        'rating': '跑赢大市',
        'target_price': '40.94',
        'source': 'xxx',
        'current_price': '', 'current_price_date': '', 'market_cap': '',
        'shares_in_issue': '', 'one_year_high': '', 'one_year_low': '',
        'a1_cn': '杨斌', 'a1_en': 'Bin Yang', 'a1_sfc': 'BSQ313', 'a1_email': 'harryyang@hfisec.com.hk',
        'a2_cn': '邓轶达', 'a2_en': 'Yida Deng', 'a2_sfc': 'BVF321', 'a2_email': 'dengyida@hfisec.com.hk',
        'a3_cn': '', 'a3_en': '', 'a3_sfc': '', 'a3_email': '',
        'a4_cn': '', 'a4_en': '', 'a4_sfc': '', 'a4_email': '',
        'a5_cn': '', 'a5_en': '', 'a5_sfc': '', 'a5_email': '',
        'a6_cn': '', 'a6_en': '', 'a6_sfc': '', 'a6_email': '',
        'a7_cn': '', 'a7_en': '', 'a7_sfc': '', 'a7_email': '',
        'a8_cn': '', 'a8_en': '', 'a8_sfc': '', 'a8_email': '',
        'abstract': 'abstract.docx', 'body': 'body.docx', 'english_summary': 'english_summary.docx',
        'company_en': '',
        'relative_chart': '', 'relative_chart_generated': False,
        'recommendation_chart': '', 'recommendation_chart_generated': False,
    }
    with open(f'{OUTPUT_DIR}/elements.yaml', 'w') as f:
        yaml.dump(elements, f, allow_unicode=True, default_flow_style=False, sort_keys=False)

    check('elements.yaml created', Path(f'{OUTPUT_DIR}/elements.yaml').exists(), True, Path(f'{OUTPUT_DIR}/elements.yaml').exists())

    return elements

# === Phase 2: Render ===

def phase2_render(elements):
    print('\n=== Phase 2: Rendering ===')

    W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'

    # Render template
    tpl = DocxTemplate(TEMPLATE_PATH)
    tpl.render(elements)
    rendered_path = f'{OUTPUT_DIR}/rendered_template.docx'
    tpl.save(rendered_path)
    print('  Template rendered')

    # Insert abstract into table cell
    doc = Document(rendered_path)
    abs_doc = Document(f'{OUTPUT_DIR}/abstract.docx')

    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                if '{{ abstract }}' in cell.text or '(Please see APPENDIX 1' in cell.text:
                    for p in list(cell._tc.findall(f'{{{W}}}p')):
                        cell._tc.remove(p)
                    for src_p in abs_doc.paragraphs:
                        new_p = deepcopy(src_p._p)
                        cell._tc.append(new_p)
    doc.save(rendered_path)
    print('  Abstract inserted')

    # Compose with docxcompose
    master = Document(rendered_path)
    composer = Composer(master)
    body_doc = Document(f'{OUTPUT_DIR}/body.docx')
    composer.append(body_doc, remove_property_fields=False)
    eng_doc = Document(f'{OUTPUT_DIR}/english_summary.docx')
    composer.append(eng_doc, remove_property_fields=False)
    output_path = f'{OUTPUT_DIR}/test_report_output.docx'
    composer.save(output_path)
    print(f'  Output saved: {output_path}')

    # Verify
    doc = Document(output_path)
    placeholders = []
    for p in doc.paragraphs:
        if re.search(r'\{\{.*?\}\}', p.text):
            placeholders.append(p.text)
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                if re.search(r'\{\{.*?\}\}', cell.text):
                    placeholders.append(cell.text[:50])

    check('No placeholders', len(placeholders) == 0, 0, len(placeholders))
    check('Final paragraphs', in_range(len(doc.paragraphs), EXPECTED['final_paragraphs']), EXPECTED['final_paragraphs'], len(doc.paragraphs))
    check('Final tables', in_range(len(doc.tables), EXPECTED['final_tables']), EXPECTED['final_tables'], len(doc.tables))

    # Cleanup temp files
    for f in Path(OUTPUT_DIR).glob('temp_*.docx'):
        f.unlink()
    for f in Path(OUTPUT_DIR).glob('rendered_*.docx'):
        f.unlink()

    return output_path

# === Main ===

def main():
    import argparse
    parser = argparse.ArgumentParser()
    parser.add_argument('--skip-extract', action='store_true')
    parser.add_argument('--skip-render', action='store_true')
    args = parser.parse_args()

    print('='*50)
    print('Report Maker E2E Regression Test')
    print('='*50)

    if args.skip_extract:
        print('\n[SKIP] Phase 1 extraction')
    else:
        phase1_extract()

    if args.skip_render:
        print('\n[SKIP] Phase 2 rendering')
    else:
        elements = yaml.safe_load(open(f'{OUTPUT_DIR}/elements.yaml'))
        phase2_render(elements)

    print('\n=== Test Complete ===')
    print(f'Output: {OUTPUT_DIR}')

if __name__ == '__main__':
    main()
