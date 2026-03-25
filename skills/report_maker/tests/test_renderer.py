"""测试 renderer.py 模块的函数。"""
import pytest
from pathlib import Path
import tempfile

import sys
sys.path.insert(0, str(Path(__file__).parent.parent / "scripts"))

from renderer import (
    prepare_render_context,
    find_placeholder_in_document,
    insert_fragment_at_placeholder,
    render_with_fragments,
    validate_template,
    find_libreoffice,
    verify_pdf_output,
    convert_to_pdf_with_retry
)


class TestPrepareRenderContext:
    """测试 prepare_render_context 函数。"""

    def test_prepare_context(self):
        """测试：准备渲染上下文。"""
        elements = {
            "ticker": "2507.HK",
            "company": "西锐飞机",
            "rating": "OUTPERFORM"
        }

        context = prepare_render_context(elements, "/reports/test")

        assert context["ticker"] == "2507.HK"
        assert context["company"] == "西锐飞机"
        assert context["rating"] == "OUTPERFORM"

    def test_add_file_paths(self):
        """测试：添加文件路径。"""
        elements = {"ticker": "2507.HK"}

        context = prepare_render_context(elements, "/reports/test")

        assert "_abstract_file" in context
        assert "_body_file" in context
        assert "abstract.docx" in context["_abstract_file"]
        assert "body.docx" in context["_body_file"]


class TestFindPlaceholderInDocument:
    """测试 find_placeholder_in_document 函数。"""

    @pytest.fixture
    def sample_doc(self, tmp_path):
        """创建包含占位符的示例文档。"""
        try:
            from docx import Document

            doc = Document()

            # 添加包含占位符的段落
            doc.add_paragraph("这是标题")
            doc.add_paragraph("{{abstract}}")
            doc.add_paragraph("其他内容")
            doc.add_paragraph("{{body}}")

            # 添加包含占位符的表格
            table = doc.add_table(rows=2, cols=2)
            table.rows[0].cells[0].text = "摘要"
            table.rows[0].cells[1].text = "{{abstract}}"

            doc_path = tmp_path / "sample.docx"
            doc.save(str(doc_path))

            return str(doc_path)
        except ImportError:
            pytest.skip("python-docx not installed")

    def test_find_placeholder_in_paragraphs(self, sample_doc):
        """测试：在段落中查找占位符。"""
        result = find_placeholder_in_document(sample_doc, "{{abstract}}")

        assert len(result) > 0
        assert any(loc["type"] == "paragraph" for loc in result)
        assert result[0]["placeholder"] == "{{abstract}}"

    def test_find_multiple_placeholders(self, sample_doc):
        """测试：查找多个占位符。"""
        result = find_placeholder_in_document(sample_doc, "{{abstract}}")

        # 应该找到多个匹配
        assert len(result) >= 2

    def test_find_nonexistent_placeholder(self, sample_doc):
        """测试：查找不存在的占位符。"""
        result = find_placeholder_in_document(sample_doc, "{{nonexistent}}")

        assert len(result) == 0


class TestInsertFragmentAtPlaceholder:
    """测试 insert_fragment_at_placeholder 函数。"""

    @pytest.fixture
    def target_doc(self, tmp_path):
        """创建目标文档。"""
        try:
            from docx import Document

            doc = Document()
            doc.add_paragraph("文档开始")
            doc.add_paragraph("{{abstract}}")
            doc.add_paragraph("文档继续")
            doc.add_paragraph("{{body}}")
            doc.add_paragraph("文档结束")

            doc_path = tmp_path / "target.docx"
            doc.save(str(doc_path))

            return str(doc_path)
        except ImportError:
            pytest.skip("python-docx not installed")

    @pytest.fixture
    def source_doc(self, tmp_path):
        """创建源文档。"""
        try:
            from docx import Document

            doc = Document()
            doc.add_paragraph("这是摘要内容")
            doc.add_paragraph("更多摘要内容")

            doc_path = tmp_path / "source.docx"
            doc.save(str(doc_path))

            return str(doc_path)
        except ImportError:
            pytest.skip("python-docx not installed")

    def test_insert_at_paragraph_placeholder(self, target_doc, source_doc, tmp_path):
        """测试：在段落占位符处插入。"""
        try:
            from docx import Document

            target = Document(target_doc)
            source = Document(source_doc)

            result = insert_fragment_at_placeholder(
                target, source, "{{abstract}}", replace_placeholder=True
            )

            assert result is True

            # 保存并验证
            output_path = tmp_path / "output.docx"
            target.save(str(output_path))
            assert Path(output_path).exists()

        except ImportError:
            pytest.skip("python-docx not installed")

    def test_insert_without_match(self, target_doc, source_doc, tmp_path):
        """测试：没有匹配占位符时的行为。"""
        try:
            from docx import Document

            target = Document(target_doc)
            source = Document(source_doc)

            result = insert_fragment_at_placeholder(
                target, source, "{{nonexistent}}", replace_placeholder=True
            )

            assert result is False

        except ImportError:
            pytest.skip("python-docx not installed")


class TestRenderWithFragments:
    """测试 render_with_fragments 函数。"""

    @pytest.fixture
    def template_doc(self, tmp_path):
        """创建模板文档。"""
        try:
            from docx import Document
            from docxtpl import DocxTemplate

            doc = Document()
            doc.add_paragraph("公司报告")
            doc.add_paragraph("公司名称: {{company}}")
            doc.add_paragraph("评级: {{rating}}")
            doc.add_paragraph("{{abstract}}")
            doc.add_paragraph("{{body}}")

            template_path = tmp_path / "template.docx"
            doc.save(str(template_path))

            return str(template_path)
        except ImportError:
            pytest.skip("python-docx not installed")

    @pytest.fixture
    def abstract_doc(self, tmp_path):
        """创建摘要文档。"""
        try:
            from docx import Document

            doc = Document()
            doc.add_paragraph("投资亮点")
            doc.add_paragraph("这是摘要内容")

            abstract_path = tmp_path / "abstract.docx"
            doc.save(str(abstract_path))

            return str(abstract_path)
        except ImportError:
            pytest.skip("python-docx not installed")

    @pytest.fixture
    def body_doc(self, tmp_path):
        """创建正文文档。"""
        try:
            from docx import Document

            doc = Document()
            doc.add_paragraph("正文开始")
            doc.add_paragraph("详细分析内容")

            body_path = tmp_path / "body.docx"
            doc.save(str(body_path))

            return str(body_path)
        except ImportError:
            pytest.skip("python-docx not installed")

    def test_render_with_fragments(self, template_doc, abstract_doc, body_doc, tmp_path):
        """测试：完整渲染流程。"""
        elements = {
            "company": "西锐飞机",
            "rating": "OUTPERFORM"
        }

        output_path = tmp_path / "output.docx"

        result = render_with_fragments(
            template_doc,
            elements,
            abstract_doc,
            body_doc,
            str(output_path)
        )

        assert result is True
        assert Path(output_path).exists()

    def test_render_without_abstract(self, template_doc, body_doc, tmp_path):
        """测试：没有摘要文档的渲染。"""
        elements = {
            "company": "西锐飞机",
            "rating": "OUTPERFORM"
        }

        output_path = tmp_path / "output.docx"
        nonexistent_abstract = tmp_path / "nonexistent_abstract.docx"

        result = render_with_fragments(
            template_doc,
            elements,
            str(nonexistent_abstract),
            body_doc,
            str(output_path)
        )

        # 应该成功（摘要文件不存在，但跳过）
        assert result is True


class TestValidateTemplate:
    """测试 validate_template 函数。"""

    @pytest.fixture
    def valid_template(self, tmp_path):
        """创建有效模板。"""
        try:
            from docx import Document

            doc = Document()
            doc.add_paragraph("模板标题")
            doc.add_paragraph("{{company}}")

            template_path = tmp_path / "template.docx"
            doc.save(str(template_path))

            return str(template_path)
        except ImportError:
            pytest.skip("python-docx not installed")

    def test_validate_valid_template(self, valid_template):
        """测试：验证有效模板。"""
        result = validate_template(valid_template)

        assert result["valid"] is True
        assert "paragraphs" in result
        assert "tables" in result
        assert "sections" in result

    def test_validate_nonexistent_template(self, tmp_path):
        """测试：验证不存在的模板。"""
        result = validate_template(str(tmp_path / "nonexistent.docx"))

        assert result["valid"] is False
        assert "error" in result


class TestFindLibreoffice:
    """测试 find_libreoffice 函数。"""

    def test_find_libreoffice_or_none(self):
        """测试：查找 LibreOffice 或返回 None。"""
        result = find_libreoffice()

        # 结果应该是路径字符串或 None
        assert result is None or isinstance(result, str)

        # 如果找到了，应该是一个有效路径
        if result:
            print(f"找到 LibreOffice: {result}")


class TestVerifyPdfOutput:
    """测试 verify_pdf_output 函数。"""

    def test_verify_nonexistent_pdf(self, tmp_path):
        """测试：验证不存在的 PDF。"""
        result = verify_pdf_output(str(tmp_path / "nonexistent.pdf"))

        assert result is False

    def test_verify_empty_pdf(self, tmp_path):
        """测试：验证空的 PDF。"""
        empty_pdf = tmp_path / "empty.pdf"
        empty_pdf.write_text("")

        result = verify_pdf_output(str(empty_pdf))

        assert result is False

    def test_verify_invalid_pdf(self, tmp_path):
        """测试：验证无效的 PDF（不是 PDF 格式）。"""
        invalid_pdf = tmp_path / "invalid.pdf"
        # 写入一些不是 PDF 的内容
        invalid_pdf.write_bytes(b"This is not a PDF")

        result = verify_pdf_output(str(invalid_pdf))

        assert result is False

    def test_verify_min_size(self, tmp_path):
        """测试：验证最小文件大小。"""
        # 创建一个有效的 PDF 头但文件很小
        small_pdf = tmp_path / "small.pdf"
        small_pdf.write_bytes(b'%PDF-1.4')

        # 使用更大的最小值
        result = verify_pdf_output(str(small_pdf), min_size=100)

        assert result is False


class TestConvertToPdfWithRetry:
    """测试 convert_to_pdf_with_retry 函数。"""

    @pytest.fixture
    def sample_docx(self, tmp_path):
        """创建示例 Word 文档。"""
        try:
            from docx import Document

            doc = Document()
            doc.add_paragraph("测试文档")
            doc.add_paragraph("这是用于 PDF 转换测试的文档")

            docx_path = tmp_path / "test.docx"
            doc.save(str(docx_path))

            return str(docx_path)
        except ImportError:
            pytest.skip("python-docx not installed")

    def test_convert_without_libreoffice(self, sample_docx, tmp_path):
        """测试：没有 LibreOffice 时的行为。"""
        # 这个测试在大多数环境中会跳过 PDF 转换
        # 因为 LibreOffice 可能未安装
        pdf_path = tmp_path / "output.pdf"

        # 注意：这个测试需要 LibreOffice 才能真正测试转换
        # 在 CI 环境中通常会跳过
        result = convert_to_pdf_with_retry(sample_docx, str(pdf_path), max_retries=1)

        # 如果没有 LibreOffice，应该返回 False
        if not find_libreoffice():
            assert result is False
        # 如果有 LibreOffice，可能成功也可能失败（取决于环境）
        # 我们只验证函数不会崩溃
        else:
            assert isinstance(result, bool)

    def test_convert_with_invalid_docx(self, tmp_path):
        """测试：转换不存在的文档。"""
        nonexistent_docx = tmp_path / "nonexistent.docx"
        pdf_path = tmp_path / "output.pdf"

        result = convert_to_pdf_with_retry(
            str(nonexistent_docx),
            str(pdf_path),
            max_retries=1
        )

        # 应该失败（文档不存在）
        assert result is False

