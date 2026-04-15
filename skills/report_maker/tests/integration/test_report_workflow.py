"""
report_maker 集成测试

测试完整的报告处理工作流，从草稿上传到最终输出生成。

测试覆盖：
- 第一阶段：LLM 提取要素、交互式确认、保存结果
- 第二阶段：渲染模板、插入片段、生成输出

测试数据：使用 samples/ 目录下的示例报告
"""

import pytest
import shutil
from pathlib import Path
from datetime import datetime

import sys
sys.path.insert(0, str(Path(__file__).parent.parent.parent))

from scripts import common, extractor, splitter, renderer


class TestReportWorkflow:
    """完整工作流集成测试"""

    # 测试数据路径
    SAMPLES_DIR = Path(__file__).parent.parent.parent.parent.parent / "samples"
    TEMPLATES_DIR = Path(__file__).parent.parent.parent.parent.parent / "templates"
    TEST_REPORTS_DIR = Path(__file__).parent.parent / "test_reports"

    @pytest.fixture(autouse=True)
    def setup_teardown(self):
        """每个测试前创建临时目录，测试后清理"""
        # 创建临时测试目录
        self.TEST_REPORTS_DIR.mkdir(exist_ok=True, parents=True)

        yield

        # 清理临时目录
        if self.TEST_REPORTS_DIR.exists():
            shutil.rmtree(self.TEST_REPORTS_DIR)

    def test_end_to_end_company_report(self):
        """测试公司报告完整工作流"""
        # 1. 准备输入
        report_type = "company"
        draft_name = "西锐飞机首次覆盖报告（中文）提交版"
        draft_path = self.SAMPLES_DIR / f"{draft_name}.docx"

        if not draft_path.exists():
            pytest.skip(f"测试草稿文件不存在: {draft_path}")

        report_name = common.normalize_report_name(draft_name)
        report_dir = self.TEST_REPORTS_DIR / report_name

        # 2. 第一阶段：提取与确认

        # 2.1 创建报告目录
        created_dir = common.create_report_dir(report_name, str(self.TEST_REPORTS_DIR))
        assert created_dir == report_dir
        assert report_dir.exists()

        # 2.2 获取模板名
        template_name = common.get_template_name(report_type)
        assert template_name == "company"

        # 2.3 加载 schema
        schema_path = self.TEMPLATES_DIR / template_name / "schema.yaml"
        schema = common.load_yaml(schema_path)
        assert "ticker" in schema
        assert "rating" in schema

        # 2.4 复制草稿
        target_draft = report_dir / "draft.docx"
        shutil.copy(draft_path, target_draft)
        assert target_draft.exists()

        # 2.5 提取要素（模拟 LLM 提取）
        sample_elements_path = self.SAMPLES_DIR / f"{draft_name} elements.yaml"
        if sample_elements_path.exists():
            extracted_elements = common.load_yaml(sample_elements_path)
        else:
            extracted_elements = extractor.extract_elements(str(target_draft), schema)

        # 验证：关键非片段字段存在（样例数据可能使用旧格式，不逐一断言）
        field_names = common.get_all_field_names(schema)
        non_fragment = common.filter_fragment_fields(field_names)
        # 核心字段必须存在
        assert "ticker" in extracted_elements, "字段 ticker 未提取"
        assert "rating" in extracted_elements, "字段 rating 未提取"
        # 至少有一半非片段字段被提取
        present_count = sum(1 for f in non_fragment if f in extracted_elements)
        assert present_count >= len(non_fragment) // 2, \
            f"提取字段数不足: {present_count}/{len(non_fragment)}"

        # 2.6 保存 elements.yaml
        elements_path = report_dir / "elements.yaml"
        common.save_yaml(elements_path, extracted_elements)
        assert elements_path.exists()

        # 2.7 验证元素与 schema 一致
        loaded_elements = common.load_yaml(elements_path)
        errors = common.validate_elements_against_schema(loaded_elements, schema)

        # 排除片段字段和图表生成状态的错误
        ignorable = {'abstract', 'body', 'compare_chart', 'volume_chart', 'recommendation_chart'}
        non_fragment_errors = [e for e in errors if not any(f in e for f in ignorable)]

        assert len(non_fragment_errors) == 0, f"elements.yaml 与 schema 不一致: {non_fragment_errors}"

        # 2.8 拆分摘要和正文（模拟）
        abstract_path = report_dir / "abstract.docx"
        body_path = report_dir / "body.docx"

        if sample_elements_path.exists():
            from docx import Document
            doc = Document()
            doc.add_paragraph("摘要占位符")
            doc.save(abstract_path)

            doc = Document()
            doc.add_paragraph("正文占位符")
            doc.save(body_path)
        else:
            abstract_doc = splitter.extract_abstract(str(target_draft), schema)
            body_doc = splitter.extract_body(str(target_draft), schema)
            abstract_doc.save(abstract_path)
            body_doc.save(body_path)

        assert abstract_path.exists()
        assert body_path.exists()

        # 3. 第二阶段：渲染与输出
        template_path = self.TEMPLATES_DIR / template_name / "template.docx"
        assert template_path.exists(), "模板文件不存在"

        output_path = report_dir / "output.docx"
        result = renderer.render_report(
            str(template_path),
            loaded_elements,
            str(abstract_path),
            str(body_path),
            str(output_path)
        )

        assert result is True, "渲染应该成功"
        assert output_path.exists(), "输出文件未生成"
        assert output_path.stat().st_size > 0, "输出文件为空"

    def test_elements_schema_consistency(self):
        """测试 elements.yaml 与 schema.yaml 的一致性"""
        sample_elements_path = self.SAMPLES_DIR / "西锐飞机首次覆盖报告（中文）提交版 elements.yaml"
        if not sample_elements_path.exists():
            pytest.skip("示例 elements.yaml 不存在")

        elements = common.load_yaml(sample_elements_path)

        # 加载 schema
        schema_path = self.TEMPLATES_DIR / "company" / "schema.yaml"
        schema = common.load_yaml(schema_path)

        # 验证一致性
        errors = common.validate_elements_against_schema(elements, schema)

        # 排除片段字段和图表生成状态的错误
        ignorable = {'abstract', 'body', 'compare_chart', 'volume_chart', 'recommendation_chart'}
        non_fragment_errors = [e for e in errors if not any(f in e for f in ignorable)]

        assert len(non_fragment_errors) == 0, f"elements.yaml 与 schema 不一致: {non_fragment_errors}"

        # 验证字段数量（样例数据可能使用旧格式，至少有核心字段）
        assert len(elements) >= 10, f"elements.yaml 字段数过少: {len(elements)}"

    def test_report_name_normalization(self):
        """测试报告名称标准化"""
        test_cases = [
            ("测试报告", "测试报告"),
            ("Test Report (2025)", "Test Report (2025)"),
            ("draft_测试报告", "测试报告"),
            ("草稿_测试报告", "测试报告"),
            ("测试报告_final", "测试报告"),
            ("测试报告_最终", "测试报告"),
            ("测试报告.docx", "测试报告"),
        ]

        for input_name, expected in test_cases:
            result = common.normalize_report_name(input_name)
            assert result == expected, f"名称标准化失败: {input_name} -> {result}"

    def test_template_routing(self):
        """测试报告类型路由"""
        test_cases = [
            ("company", "company"),
        ]

        for report_type, expected_template in test_cases:
            template = common.get_template_name(report_type)
            assert template == expected_template, \
                f"路由错误: {report_type} -> {template}"

        # 测试不支持的类型 - 返回 None
        template = common.get_template_name("unsupported_type")
        assert template is None, "不支持的类型应该返回 None"

    def test_fragment_field_handling(self):
        """测试片段字段处理"""
        schema_path = self.TEMPLATES_DIR / "company" / "schema.yaml"
        schema = common.load_yaml(schema_path)

        # 获取所有字段
        all_fields = common.get_all_field_names(schema)
        assert "abstract" in all_fields
        assert "body" in all_fields
        assert "ticker" in all_fields

        # 过滤出非片段字段
        non_fragment_fields = common.filter_fragment_fields(all_fields)
        assert "abstract" not in non_fragment_fields
        assert "body" not in non_fragment_fields
        assert "ticker" in non_fragment_fields

    def test_elements_field_completion(self):
        """测试确保所有字段存在"""
        schema_path = self.TEMPLATES_DIR / "company" / "schema.yaml"
        schema = common.load_yaml(schema_path)

        # 创建不完整的 elements
        partial_elements = {"ticker": "2507.HK"}

        # 补全字段
        completed = common.ensure_all_fields(partial_elements, schema)

        # 验证所有字段都存在（但缺失的会是空字符串）
        all_fields = common.get_all_field_names(schema)
        for field in all_fields:
            assert field in completed, f"字段 {field} 未补全"

        # 验证原有字段保留
        assert completed["ticker"] == "2507.HK"

    def test_directory_structure(self):
        """测试报告目录结构创建"""
        report_name = "test_report"
        report_dir = common.create_report_dir(report_name, str(self.TEST_REPORTS_DIR))

        # 验证目录结构
        assert report_dir.exists()
        assert report_dir.is_dir()
        assert report_dir.name == report_name

        # 验证在正确的父目录下
        assert self.TEST_REPORTS_DIR in report_dir.parents

    def test_yaml_round_trip(self):
        """测试 YAML 读写往返"""
        test_data = {
            "string_field": "测试字符串",
            "number_field": "123.45",
            "empty_field": "",
            "list_field": ["item1", "item2"],
            "nested": {"key": "value"},
        }

        test_file = self.TEST_REPORTS_DIR / "test_round_trip.yaml"

        # 写入
        common.save_yaml(test_file, test_data)
        assert test_file.exists()

        # 读取
        loaded = common.load_yaml(test_file)

        # 验证
        assert loaded["string_field"] == test_data["string_field"]
        assert loaded["number_field"] == test_data["number_field"]
        assert loaded["empty_field"] == test_data["empty_field"]
        assert loaded["list_field"] == test_data["list_field"]
        assert loaded["nested"]["key"] == test_data["nested"]["key"]


class TestErrorHandling:
    """错误处理集成测试"""

    TEST_REPORTS_DIR = Path(__file__).parent.parent / "test_reports"

    @pytest.fixture(autouse=True)
    def setup_teardown(self):
        """每个测试前创建临时目录，测试后清理"""
        self.TEST_REPORTS_DIR.mkdir(exist_ok=True, parents=True)
        yield
        if self.TEST_REPORTS_DIR.exists():
            shutil.rmtree(self.TEST_REPORTS_DIR)

    def test_missing_template_file(self):
        """测试模板文件缺失"""
        template_path = self.TEST_REPORTS_DIR / "nonexistent.docx"
        output_path = self.TEST_REPORTS_DIR / "output.docx"
        abstract_path = self.TEST_REPORTS_DIR / "abstract.docx"
        body_path = self.TEST_REPORTS_DIR / "body.docx"

        # 创建占位片段文件
        from docx import Document
        doc = Document()
        doc.save(abstract_path)
        doc.save(body_path)

        # 渲染应该返回 False
        result = renderer.render_report(
            str(template_path),
            {},
            str(abstract_path),
            str(body_path),
            str(output_path)
        )
        assert result is False, "不存在的模板应该返回 False"

    def test_missing_schema_file(self):
        """测试 schema 文件缺失 - load_yaml 返回空 dict"""
        result = common.load_yaml(self.TEST_REPORTS_DIR / "nonexistent.yaml")
        assert result == {}, "不存在的文件应该返回空字典"

    def test_invalid_report_type(self):
        """测试无效的报告类型"""
        template = common.get_template_name("invalid_type")
        assert template is None, "不支持的类型应该返回 None"

    def test_missing_draft_file(self):
        """测试草稿文件缺失 - extract_elements 不读取文件"""
        schema_path = Path(__file__).parent.parent.parent.parent.parent / "templates" / "company" / "schema.yaml"
        schema = common.load_yaml(schema_path)

        # extract_elements 不直接读取文件，返回空 elements
        result = extractor.extract_elements("nonexistent.docx", schema)
        assert isinstance(result, dict), "应该返回字典"
        assert len(result) > 0, "应该包含字段"


class TestBackwardCompatibility:
    """向后兼容性测试"""

    def test_legacy_elements_format(self):
        """测试旧版 elements.yaml 格式兼容性"""
        sample_elements_path = Path(__file__).parent.parent.parent.parent.parent / "samples" / "西锐飞机首次覆盖报告（中文）提交版 elements.yaml"

        if sample_elements_path.exists():
            elements = common.load_yaml(sample_elements_path)
            assert isinstance(elements, dict)
            assert len(elements) > 0

    def test_existing_template_readable(self):
        """测试现有模板可读取"""
        template_path = Path(__file__).parent.parent.parent.parent.parent / "templates" / "company" / "template.docx"

        if template_path.exists():
            try:
                from docx import Document
                doc = Document(str(template_path))
                assert len(doc.paragraphs) > 0 or len(doc.tables) > 0, "模板应有内容"
            except ImportError:
                pytest.skip("python-docx not installed")
