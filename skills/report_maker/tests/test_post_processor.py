"""test_post_processor.py - post_processor.py 模块测试。

测试后处理功能：验证附录、字体约束、分析师认证扩展等。
"""

import pytest
from pathlib import Path
import zipfile
from unittest.mock import patch

import sys
sys.path.insert(0, str(Path(__file__).parent.parent / "scripts"))


class TestVerifyAppendixWithLxml:
    """测试 _verify_appendix_with_lxml 函数。"""

    def test_returns_not_found_for_empty_document(self, tmp_path):
        """测试：空文档返回未找到。"""
        from post_processor import _verify_appendix_with_lxml

        try:
            from docx import Document
        except ImportError:
            pytest.skip("python-docx not installed")

        doc = Document()
        doc_path = tmp_path / "empty.docx"
        doc.save(str(doc_path))

        result = _verify_appendix_with_lxml(str(doc_path))

        assert result["found"] is False

    def test_finds_appendix_in_paragraph(self, tmp_path):
        """测试：在段落中找到附录标记。"""
        from post_processor import _verify_appendix_with_lxml

        try:
            from docx import Document
        except ImportError:
            pytest.skip("python-docx not installed")

        doc = Document()
        doc.add_paragraph("Some content")
        doc.add_paragraph("附录APPENDIX 1")

        doc_path = tmp_path / "with_appendix.docx"
        doc.save(str(doc_path))

        result = _verify_appendix_with_lxml(str(doc_path))

        assert result["found"] is True
        assert "附录" in result["text"]

    def test_finds_important_disclosure(self, tmp_path):
        """测试：找到 IMPORTANT DISCLOSURE。"""
        from post_processor import _verify_appendix_with_lxml

        try:
            from docx import Document
        except ImportError:
            pytest.skip("python-docx not installed")

        doc = Document()
        doc.add_paragraph("IMPORTANT DISCLOSURE")
        doc.add_paragraph("Some legal text")

        doc_path = tmp_path / "with_disclosure.docx"
        doc.save(str(doc_path))

        result = _verify_appendix_with_lxml(str(doc_path))

        assert result["found"] is True

    def test_finds_appendix_in_table(self, tmp_path):
        """测试：在表格中找到附录标记。"""
        from post_processor import _verify_appendix_with_lxml

        try:
            from docx import Document
        except ImportError:
            pytest.skip("python-docx not installed")

        doc = Document()
        table = doc.add_table(rows=1, cols=1)
        table.rows[0].cells[0].text = "附录APPENDIX content"

        doc_path = tmp_path / "with_table_appendix.docx"
        doc.save(str(doc_path))

        result = _verify_appendix_with_lxml(str(doc_path))

        assert result["found"] is True

    def test_handles_nonexistent_file(self):
        """测试：处理不存在的文件。"""
        from post_processor import _verify_appendix_with_lxml

        result = _verify_appendix_with_lxml("nonexistent.docx")

        assert result["found"] is False
        assert "error" in result


class TestApplyFontConstraints:
    """测试 _apply_font_constraints 函数。"""

    def test_returns_false_when_docx_unavailable(self):
        """测试：docx 不可用时返回 False。"""
        from post_processor import _apply_font_constraints

        with patch('post_processor.DOCX_AVAILABLE', False):
            result = _apply_font_constraints("any_path.docx")

            assert result is False

    def test_returns_false_for_nonexistent_file(self):
        """测试：文件不存在时返回 False。"""
        from post_processor import _apply_font_constraints

        with patch('post_processor.DOCX_AVAILABLE', True):
            result = _apply_font_constraints("nonexistent.docx")

            assert result is False


class TestExpandAnalystCertification:
    """测试 expand_analyst_certification 函数。"""

    def test_returns_zero_when_docx_unavailable(self):
        """测试：docx 不可用时返回 0。"""
        from post_processor import expand_analyst_certification

        with patch('post_processor.DOCX_AVAILABLE', False):
            result = expand_analyst_certification("any_path.docx", {})

            assert result == 0

    def test_returns_zero_for_single_analyst(self, tmp_path):
        """测试：只有一个分析师时返回 0。"""
        from post_processor import expand_analyst_certification

        try:
            from docx import Document
        except ImportError:
            pytest.skip("python-docx not installed")

        doc = Document()
        doc.add_paragraph("分析师：杨斌")

        doc_path = tmp_path / "single_analyst.docx"
        doc.save(str(doc_path))

        elements = {"a1_name": "杨斌"}

        result = expand_analyst_certification(str(doc_path), elements)

        assert result == 0

    def test_returns_zero_when_no_certification_paragraph(self, tmp_path):
        """测试：没有认证段落时返回 0。"""
        from post_processor import expand_analyst_certification

        try:
            from docx import Document
        except ImportError:
            pytest.skip("python-docx not installed")

        doc = Document()
        doc.add_paragraph("普通段落，没有分析师认证")

        doc_path = tmp_path / "no_certification.docx"
        doc.save(str(doc_path))

        elements = {"a1_name": "分析师A", "a2_name": "分析师B"}

        result = expand_analyst_certification(str(doc_path), elements)

        assert result == 0

    def test_returns_zero_for_nonexistent_file(self):
        """测试：文件不存在时返回 0。"""
        from post_processor import expand_analyst_certification

        with patch('post_processor.DOCX_AVAILABLE', True):
            result = expand_analyst_certification("nonexistent.docx", {})

            assert result == 0
