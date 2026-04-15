"""test_sectpr_manager.py - sectpr_manager.py 模块测试。

测试 sectPr（分节属性）管理功能。
"""

import pytest
from pathlib import Path
import zipfile
from unittest.mock import patch

import sys
sys.path.insert(0, str(Path(__file__).parent.parent / "scripts"))


class TestSaveTemplateFinalSectPr:
    """测试 _save_template_final_sectPr 函数。"""

    def test_returns_none_for_nonexistent_file(self):
        """测试：文件不存在时返回 None。"""
        from sectpr_manager import _save_template_final_sectPr

        result = _save_template_final_sectPr("nonexistent.docx")

        assert result is None


class TestSaveAllTemplateSectPrs:
    """测试 _save_all_template_sectPrs 函数。"""

    def test_returns_none_for_nonexistent_file(self):
        """测试：文件不存在时返回 None。"""
        from sectpr_manager import _save_all_template_sectPrs

        result = _save_all_template_sectPrs("nonexistent.docx")

        assert result is None

    def test_handles_file_without_body(self, tmp_path):
        """测试：处理没有 body 的文件。"""
        from sectpr_manager import _save_all_template_sectPrs

        doc_path = tmp_path / "no_body.docx"
        with zipfile.ZipFile(str(doc_path), 'w') as z:
            z.writestr("[Content_Types].xml", b"<Types/>")
            z.writestr("word/document.xml", b"<document/>")

        result = _save_all_template_sectPrs(str(doc_path))

        assert result is None


class TestRestoreAllSectPrs:
    """测试 _restore_all_sectPrs 函数。"""

    def test_returns_false_for_nonexistent_file(self):
        """测试：文件不存在时返回 False。"""
        from sectpr_manager import _restore_all_sectPrs

        result = _restore_all_sectPrs("nonexistent.docx", {})

        assert result is False

    def test_returns_false_for_file_without_document_xml(self, tmp_path):
        """测试：文件没有 document.xml 时返回 False。"""
        from sectpr_manager import _restore_all_sectPrs

        doc_path = tmp_path / "no_docxml.docx"
        with zipfile.ZipFile(str(doc_path), 'w') as z:
            z.writestr("[Content_Types].xml", b"<Types/>")

        result = _restore_all_sectPrs(str(doc_path), {})

        assert result is False

    def test_handles_empty_template_sectprs(self, tmp_path):
        """测试：处理空的 template_sectPrs。"""
        from sectpr_manager import _restore_all_sectPrs

        try:
            from docx import Document
        except ImportError:
            pytest.skip("python-docx not installed")

        doc = Document()
        doc.add_paragraph("Test content")

        doc_path = tmp_path / "test.docx"
        doc.save(str(doc_path))

        result = _restore_all_sectPrs(str(doc_path), {})

        # 空 sectPrs 不应该导致崩溃，但可能返回 False
        assert isinstance(result, bool)


class TestRestoreFinalSectPr:
    """测试 _restore_final_sectPr 函数。"""

    def test_returns_false_for_nonexistent_file(self):
        """测试：文件不存在时返回 False。"""
        from sectpr_manager import _restore_final_sectPr

        result = _restore_final_sectPr("nonexistent.docx", None)

        assert result is False

    def test_returns_false_for_file_without_document_xml(self, tmp_path):
        """测试：文件没有 document.xml 时返回 False。"""
        from sectpr_manager import _restore_final_sectPr

        doc_path = tmp_path / "no_docxml.docx"
        with zipfile.ZipFile(str(doc_path), 'w') as z:
            z.writestr("[Content_Types].xml", b"<Types/>")

        result = _restore_final_sectPr(str(doc_path), None)

        assert result is False


class TestGetTemplateEnglishSummaryMargins:
    """测试 _get_template_english_summary_margins 函数。"""

    def test_returns_empty_dict_for_nonexistent_file(self):
        """测试：文件不存在时返回空字典。"""
        from sectpr_manager import _get_template_english_summary_margins

        result = _get_template_english_summary_margins("nonexistent.docx")

        assert result == {}

    def test_returns_empty_dict_when_marker_not_found(self, tmp_path):
        """测试：找不到 APPENDIX 1 标记时返回空字典。"""
        from sectpr_manager import _get_template_english_summary_margins

        try:
            from docx import Document
        except ImportError:
            pytest.skip("python-docx not installed")

        doc = Document()
        doc.add_paragraph("No APPENDIX 1 here")

        doc_path = tmp_path / "no_marker.docx"
        doc.save(str(doc_path))

        result = _get_template_english_summary_margins(str(doc_path))

        assert result == {}
